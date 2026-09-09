#!/usr/bin/env python3
"""Replace only Part 4 of the 2026-09-08 meeting report.

The source DOCX is treated as an OPC/ZIP package.  A temporary fragment is
created with python-docx, then its body XML and image relationships are merged
into the original package between the Part 4 and Part 5 headings.
"""
from __future__ import annotations

import copy
import csv
import shutil
import tempfile
import zipfile
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "신작물보호제_미팅보고서_260908_한서윤.docx"
OUTPUT = ROOT / "신작물보호제_미팅보고서_260908_한서윤_part4수정.docx"
ASSET_DIR = ROOT / "04_모델산출물" / "회의보고서_part4_시각화"
ARM_CSV = ROOT / "04_모델산출물" / "v6_jurisdiction" / "감작_CTarm_격자.csv"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W_NS, "r": R_NS}

BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GRAY = "F2F2F2"
TEXT = "222222"
DOC_FONT = "Apple SD Gothic Neo"


def configure_plotting() -> None:
    mpl.rcParams.update({
        "font.family": "Apple SD Gothic Neo",
        "axes.unicode_minus": False,
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "axes.edgecolor": "#777777",
        "axes.labelcolor": "#222222",
        "xtick.color": "#333333",
        "ytick.color": "#333333",
        "text.color": "#222222",
        "font.size": 10,
    })


def save_fig(fig: plt.Figure, path: Path) -> None:
    fig.savefig(path, dpi=300, facecolor="white", bbox_inches=None)
    plt.close(fig)


def make_b4_figure(path: Path) -> None:
    labels = ["피부", "감작"]
    counts = [128, 211]
    colors = ["#4E79A7", "#E07B39"]
    fig, ax = plt.subplots(figsize=(6.25, 2.65), layout="constrained")
    y = np.arange(len(labels))
    bars = ax.barh(y, counts, color=colors, height=0.52)
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlim(0, 230)
    ax.set_xlabel("현재 음성으로 남아 있는 negrec 행 수")
    ax.set_title("B4: SDS 침묵을 음성으로 유지하면 영향을 받는 학습 표본")
    ax.grid(axis="x", color="#D9D9D9", linewidth=0.6)
    ax.set_axisbelow(True)
    for bar, value in zip(bars, counts):
        ax.text(value + 5, bar.get_y() + bar.get_height()/2, f"{value}행",
                va="center", fontweight="bold")
    ax.axvspan(20, 30, color="#59A14F", alpha=0.18)
    ax.text(25, 1.62, "1차 원문 표본 20–30건\n(두 endpoint에서 층화 추출)",
            ha="center", va="top", color="#2E6B2E", fontsize=9)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    save_fig(fig, path)


def load_arm_data() -> dict[tuple[str, str], dict[str, float]]:
    out = {}
    with ARM_CSV.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            out[(row["결측처리"], row["arm"])] = {
                "auc": float(row["ROC_AUC"]),
                "sd": float(row["ROC_AUC_sd"]),
                "coverage": float(row["감작커버리지"]),
                "mcc": float(row["MCC"]),
            }
    return out


def make_d7_d11_figure(path: Path) -> None:
    data = load_arm_data()
    fig, (ax1, ax2) = plt.subplots(
        1, 2, figsize=(6.7, 3.05), gridspec_kw={"width_ratios": [0.78, 1.55]},
        layout="constrained")

    miss_names = ["화학 피처", "CT 26열"]
    miss = [19.6, 29.8]
    bars = ax1.bar(miss_names, miss, color=["#4E79A7", "#E07B39"], width=0.62)
    ax1.set_ylim(0, 35)
    ax1.set_ylabel("결측률 (%)")
    ax1.set_title("D7: 결측 규모")
    ax1.grid(axis="y", color="#D9D9D9", linewidth=0.6)
    ax1.set_axisbelow(True)
    for bar, value in zip(bars, miss):
        ax1.text(bar.get_x() + bar.get_width()/2, value + 1, f"{value:.1f}%",
                 ha="center", fontweight="bold")

    arm_ids = ["A2_aug_nc_unknown", "A3_aug_annexvi", "A3s_aug_annexvi_strict"]
    arm_names = ["A2", "A3", "A3s"]
    x = np.arange(3)
    for mode, color, marker, dx, label in [
        ("nan0", "#4E79A7", "o", -0.07, "0 채움 (현행 정규)"),
        ("native", "#E07B39", "s", 0.07, "네이티브 결측"),
    ]:
        vals = [data[(mode, a)]["auc"] for a in arm_ids]
        errs = [data[(mode, a)]["sd"] for a in arm_ids]
        ax2.errorbar(x + dx, vals, yerr=errs, color=color, marker=marker,
                     linestyle="-", linewidth=1.5, capsize=3, label=label)
        for xx, val in zip(x + dx, vals):
            ax2.text(xx, val + 0.0028, f"{val:.3f}", ha="center", fontsize=8)
    ax2.set_xticks(x, arm_names)
    ax2.set_ylim(0.71, 0.77)
    ax2.set_ylabel("감작 ROC-AUC (평균 ± SD)")
    ax2.set_title("D11: 결측 규약에 따라 arm 순위가 달라짐")
    ax2.grid(axis="y", color="#D9D9D9", linewidth=0.6)
    ax2.set_axisbelow(True)
    ax2.legend(frameon=False, loc="lower right", fontsize=8)
    ax2.text(0.5, 0.7115,
             "A3-A2: 0 채움 +0.0016 (2/5 seed)  |  native +0.0219 (5/5 seed)",
             ha="center", va="bottom", fontsize=8.2)
    for ax in (ax1, ax2):
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
    save_fig(fig, path)


def make_d8_figure(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(6.25, 2.15), layout="constrained")
    ax.set_xlim(0.20, 0.56)
    ax.set_ylim(0, 1)
    ax.set_yticks([])
    ax.set_xlabel("위험 판정 임계값")
    ax.set_title("D8: 관측된 참고 범위와 현재 임계값은 다르지만, 비용비가 먼저다")
    ax.axvspan(0.298, 0.400, ymin=0.34, ymax=0.74, color="#4E79A7", alpha=0.25)
    ax.hlines(0.54, 0.298, 0.400, color="#4E79A7", linewidth=7)
    ax.text(0.349, 0.79, "OOF 후향 참고 0.298–0.400\n(seed 4 단독, 성능 주장 불가)",
            ha="center", va="bottom", fontsize=9)
    ax.plot([0.5], [0.54], marker="D", color="#E15759", markersize=9)
    ax.text(0.5, 0.79, "현재 고정 0.5", ha="center", va="bottom",
            color="#9C2F31", fontweight="bold")
    ax.annotate("임계값을 낮추면\n위험 누락↓ · 오판정↑", xy=(0.255, 0.20),
                ha="center", va="center", fontsize=9,
                arrowprops={"arrowstyle": "<-", "color": "#666666"},
                xytext=(0.33, 0.20))
    ax.annotate("임계값을 높이면\n위험 누락↑ · 오판정↓", xy=(0.535, 0.20),
                ha="center", va="center", fontsize=9,
                arrowprops={"arrowstyle": "<-", "color": "#666666"},
                xytext=(0.46, 0.20))
    ax.grid(axis="x", color="#D9D9D9", linewidth=0.6)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    save_fig(fig, path)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = TEXT,
                  size: float = 8.6) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, *, size: float, bold: bool = False, color: str = TEXT) -> None:
    run.bold = bold
    run.font.name = DOC_FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), DOC_FONT)
    r_fonts.set(qn("w:hAnsi"), DOC_FONT)
    r_fonts.set(qn("w:eastAsia"), DOC_FONT)


def add_table(doc: Document, headers: list[str], rows: list[list[str]],
              widths_cm: list[float], header_fill: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, (header, width) in enumerate(zip(headers, widths_cm)):
        cell = table.rows[0].cells[i]
        cell.width = Cm(width)
        shade_cell(cell, header_fill)
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8.2)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, (value, width) in enumerate(zip(row, widths_cm)):
            cells[i].width = Cm(width)
            shade_cell(cells[i], "FFFFFF" if ridx % 2 == 0 else LIGHT_GRAY)
            set_cell_text(cells[i], value, size=8.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body(doc: Document, text: str, *, bold: bool = False, size: float = 9.2,
             after: float = 4, color: str = TEXT) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)


def add_picture(doc: Document, image: Path, caption: str, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(image), width=Cm(16.0))
    # Accessible description in drawing properties.
    for doc_pr in p._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", alt)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(6)
    r = cp.add_run(caption)
    set_run_font(r, size=8, color="666666")


def page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_fragment(path: Path, figures: dict[str, Path]) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    h1 = doc.styles["Heading 1"]
    h1.font.name = DOC_FONT
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(BLUE)
    h1._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), DOC_FONT)
    h2 = doc.styles["Heading 2"]
    h2.font.name = DOC_FONT
    h2.font.size = Pt(12)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), DOC_FONT)

    doc.add_heading("4. 회의에서 결정할 사항", level=1)
    add_body(doc,
             "이번 회의에서는 수치를 바로 확정하는 것이 아니라, 어떤 근거를 정답으로 인정하고 "
             "어떤 절차로 학습에 반영할지를 결정합니다. 아래 네 항목은 서로 독립적이지 않습니다. "
             "B4는 라벨 신뢰도, D7은 결측 처리, D11은 감작 CT 근거 범위, D8은 최종 판정 기준을 정합니다.",
             bold=True, size=9.4)
    add_table(doc,
              ["ID", "지금 결정할 질문", "현재 처리", "회의에서 남길 결과"],
              [
                  ["B4", "SDS 침묵을 음성으로 볼 것인가?", "피부·감작은 유지", "표본 판독 방식과 유지/마스크 기준"],
                  ["D7", "결측을 0으로 볼 것인가?", "0 채움 정규, native 병기", "주 분석 규약과 비교 실험 범위"],
                  ["D11", "감작 CT 근거를 어디까지 인정할 것인가?", "현행 A3", "A2/A3/A3s 중 정본 arm"],
                  ["D8", "어떤 오류를 더 비싸게 볼 것인가?", "임계값 0.5 고정", "놓침:오판정 비용비와 산정 절차"],
              ], [1.2, 5.0, 4.0, 6.0])

    doc.add_heading("4-1. B4 — SDS에 문구가 없으면 음성인가?", level=2)
    add_body(doc,
             "현재 피부 128행과 감작 211행은 SDS에서 관련 문구를 찾지 못했다는 이유로 NC에 포함되어 "
             "있습니다. 그러나 문구 부재는 ‘시험 결과 음성’, ‘시험하지 않음’, ‘기재 생략’을 구분하지 "
             "못합니다. 그대로 학습하면 미시험 사례를 안전 사례로 가르칠 수 있고, 전부 제외하면 339행을 "
             "잃습니다.")
    add_picture(doc, figures["b4"],
                "그림 6. B4의 영향 규모. 막대는 현재 negrec 행 수이며, 녹색 구간은 먼저 판독할 20–30건 표본 규모를 나타낸다.",
                "피부 negrec 128행, 감작 negrec 211행. 먼저 두 endpoint에서 층화한 20~30건을 원문 판독하도록 제안한다.")
    add_table(doc, ["선택", "장점", "위험/비용", "판정 근거"], [
        ["① 현행 유지", "학습행을 보존", "라벨 잡음이 계속될 수 있음", "현재는 실제 음성 비율을 모름"],
        ["② 20–30건 표본 후 결정 (권고)", "실제 음성 비율을 보고 결정", "짧은 원문 판독 필요", "endpoint·출처별 층화표본"],
        ["③ 즉시 마스크", "미시험=음성 위험을 즉시 차단", "피부 128행·감작 211행 손실", "눈과 달리 계통 오류 증거가 약함"],
    ], [3.5, 4.0, 4.4, 4.3], header_fill="548235")
    add_body(doc,
             "회의 결정 문장: ‘피부·감작 negrec에서 총 20–30건을 endpoint와 출처별로 층화 추출해 "
             "Section 11을 판독하고, 실제 음성 비율과 확인 불가 비율을 보고한 뒤 유지/마스크를 확정한다.’",
             bold=True, color="375623")

    page_break(doc)
    doc.add_heading("4-2. D7 + D11 — 결측 규약과 감작 CT arm은 함께 결정", level=2)
    add_body(doc,
             "D7은 단순한 전처리 선택이 아닙니다. 현재 0 채움은 미지값을 숫자 0과 같게 만들어 ‘근거 없음’을 "
             "‘위험 성분 없음’처럼 표현할 수 있습니다. 반면 native 결측으로 바꾸면 과거 성능과 직접 비교하기 "
             "어렵습니다. 이 선택은 감작 CT arm의 우열까지 바꿉니다.")
    add_picture(doc, figures["d7d11"],
                "그림 7. 왼쪽은 현행 피처 결측률, 오른쪽은 감작 CT arm별 ROC-AUC다. 오차막대는 5개 seed의 SD다.",
                "화학 피처 결측률 19.6%, CT 26열 결측률 29.8%. 0 채움에서는 A2와 A3가 거의 같지만 native 결측에서는 A3가 A2보다 0.0219 높다.")
    add_table(doc, ["arm", "실제 처리 규칙", "감작 CT 커버리지", "해석상 주의"], [
        ["A2", "추가 양성 범주는 사용하되 NC는 미지로 둠", "32.3%", "출처 부재를 음성으로 세지 않음"],
        ["A3 (현행)", "annex_vi tier만 사용하고 NC도 알려진 값으로 셈", "49.6%", "현 tier 파서가 자가신고-only CAS 일부를 Annex VI로 오인"],
        ["A3s", "메모 기반으로 Annex VI tier를 엄격 교정", "44.9%", "오인은 줄지만 0 채움에서 A2보다 낮음"],
    ], [2.0, 6.2, 3.0, 5.0], header_fill="8064A2")
    add_body(doc,
             "중요: A2를 ‘Annex VI만 사용’으로 설명하면 실제 코드와 다릅니다. A2의 핵심은 NC를 음성으로 "
             "채우지 않고 미지로 남기는 것입니다. A3는 명목상 Annex VI arm이지만 현재 tier 파서 오염이 "
             "있고, A3s가 그 교정판입니다.", bold=True, color="7030A0")
    add_body(doc,
             "수치 해석: 현행 0 채움에서 A3-A2는 +0.0016이고 5개 seed 중 2개에서만 양수입니다. "
             "native에서는 +0.0219이며 5개 seed 모두 양수입니다. 따라서 D11을 성능만으로 먼저 정하면 "
             "D7 선택을 몰래 전제하게 됩니다.")
    add_table(doc, ["결정", "권고안", "회의에서 확인할 것"], [
        ["D7", "0 채움을 정규 베이스라인으로 유지하고 native를 감도분석으로 병기", "native를 주 분석으로 승격할 조건과 과거 재산출 범위"],
        ["D11", "세 endpoint의 규칙을 맞추기 위해 A2 통일", "근거 일관성을 위해 native AUC 0.0219 차이를 감수할지"],
    ], [2.0, 7.0, 7.2])
    add_body(doc,
             "회의 결정란: D7 주 분석 = [ 0 채움 / native ] · 비교 실험 = [ 병기 / 미실시 ] · "
             "D11 감작 arm = [ A2 / A3 / A3s ]", bold=True)

    page_break(doc)
    doc.add_heading("4-3. D8 — 임계값이 아니라 오류 비용비를 먼저 결정", level=2)
    add_body(doc,
             "ROC-AUC는 표본의 순위를 얼마나 잘 가르는지를 보지만, 실제 위험/안전 판정은 임계값에 따라 "
             "달라집니다. 현재 0.5에서는 MCC가 0.237–0.417로 낮습니다. OOF에서 후향적으로 고른 참고 "
             "임계값은 0.298–0.400이었지만, 이는 마지막 seed 하나를 평가자료에서 최적화한 값이므로 성능 "
             "주장이나 최종 기준으로 사용할 수 없습니다.")
    add_picture(doc, figures["d8"],
                "그림 8. 현재 고정 임계값과 OOF 후향 참고 범위. 참고 범위는 방향만 보여주며 최종 임계값이 아니다.",
                "현재 임계값은 0.5이고 후향 참고 범위는 0.298~0.400이다. 임계값을 낮추면 위험 누락은 줄지만 오판정은 늘어난다.")
    add_table(doc, ["회의에서 정할 것", "의미", "결정 후 구현"], [
        ["놓침(FN):오판정(FP) 비용비", "위험 제품을 안전으로 놓치는 오류를 몇 배 더 무겁게 볼지", "비용함수에 반영"],
        ["비용비를 endpoint별로 둘지", "눈·피부·감작의 위해성과 운영비용이 같은지", "endpoint별 또는 공통 임계값"],
        ["허용 가능한 최소 재현율", "비용비 외에 반드시 지킬 안전 하한이 있는지", "제약조건으로 적용"],
    ], [4.5, 6.4, 5.3], header_fill="C65911")
    add_body(doc,
             "재현 가능한 산정 절차: 각 outer 학습 fold 안에서만 inner CV로 임계값을 고르고, 손대지 않은 "
             "outer 평가 fold에서 성능을 계산합니다. 최종 보고에는 ROC-AUC와 함께 선택 임계값, 재현율, "
             "특이도, MCC, 비용비를 모두 남깁니다.")
    add_body(doc,
             "회의 결정란: 놓침(FN):오판정(FP) = [      : 1 ] · 적용 = [ endpoint별 / 공통 ] · "
             "최소 재현율 = [        ]", bold=True)

    doc.add_heading("4-4. 회의 진행 순서와 기록 위치", level=2)
    add_table(doc, ["순서", "논의", "회의 종료 조건"], [
        ["1", "B4", "20–30건 표본 설계·담당·마감 확정"],
        ["2", "D7", "주 결측 규약과 감도분석 범위 확정"],
        ["3", "D11", "D7 결정을 전제로 감작 arm 확정"],
        ["4", "D8", "비용비·최소 재현율·임계값 산정 절차 확정"],
    ], [1.4, 2.6, 12.2])
    add_body(doc,
             "결정 기록: dataset_배정_260904.xlsx > 총책임자_결정_9건 시트. "
             "근거 데이터: 감작_CTarm_격자.csv, 감작_CTarm_대조.csv, 전체지표_요약.json. "
             "회의 후 결정·근거·결정일을 같은 행에 입력합니다.", size=8.4, color="666666")
    add_body(doc,
             "시각화 방법 참고: Kassis, T., Agarwal, V., He, Y., Patel, D., & Brueckner, A. M. (2026). "
             "Scientific Agent Skills: A Library of Procedural Knowledge for Research Agents. "
             "arXiv:2609.00065. https://doi.org/10.48550/arXiv.2609.00065",
             size=7.5, color="777777")
    page_break(doc)
    doc.save(path)


def paragraph_text(el: etree._Element) -> str:
    return "".join(el.xpath(".//w:t/text()", namespaces=NS)).strip()


def merge_fragment(source: Path, fragment: Path, output: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        src_dir, frag_dir = tmp / "src", tmp / "frag"
        with zipfile.ZipFile(source) as z:
            z.extractall(src_dir)
        with zipfile.ZipFile(fragment) as z:
            z.extractall(frag_dir)
        for link in src_dir.rglob("*"):
            if link.is_symlink():
                link.unlink()
        for link in frag_dir.rglob("*"):
            if link.is_symlink():
                link.unlink()

        parser = etree.XMLParser(remove_blank_text=False)
        src_tree = etree.parse(str(src_dir / "word/document.xml"), parser)
        frag_tree = etree.parse(str(frag_dir / "word/document.xml"), parser)
        src_body = src_tree.find(f".//{{{W_NS}}}body")
        frag_body = frag_tree.find(f".//{{{W_NS}}}body")

        children = list(src_body)
        start = next(i for i, el in enumerate(children)
                     if paragraph_text(el) == "4. 회의에서 결정할 사항")
        end = next(i for i, el in enumerate(children[start + 1:], start + 1)
                   if paragraph_text(el) == "5. 역할 분담")

        # Copy fragment image parts and assign source-safe relationship IDs.
        rel_path = src_dir / "word/_rels/document.xml.rels"
        frag_rel_path = frag_dir / "word/_rels/document.xml.rels"
        src_rels = etree.parse(str(rel_path), parser)
        frag_rels = etree.parse(str(frag_rel_path), parser)
        src_rel_root = src_rels.getroot()
        frag_rel_root = frag_rels.getroot()
        used_ids = {r.get("Id") for r in src_rel_root}
        next_id = max(int(x[3:]) for x in used_ids if x.startswith("rId")) + 1
        media_dir = src_dir / "word/media"
        media_dir.mkdir(exist_ok=True)
        existing_media = {p.name for p in media_dir.iterdir()}
        rel_map = {}
        for rel in frag_rel_root:
            if not rel.get("Type", "").endswith("/image"):
                continue
            old_id = rel.get("Id")
            old_target = rel.get("Target")
            src_img = frag_dir / "word" / old_target
            suffix = src_img.suffix.lower()
            idx = 1
            while f"part4_image{idx}{suffix}" in existing_media:
                idx += 1
            new_name = f"part4_image{idx}{suffix}"
            existing_media.add(new_name)
            shutil.copy2(src_img, media_dir / new_name)
            new_id = f"rId{next_id}"
            next_id += 1
            new_rel = etree.Element(f"{{{PKG_REL_NS}}}Relationship")
            new_rel.set("Id", new_id)
            new_rel.set("Type", rel.get("Type"))
            new_rel.set("Target", f"media/{new_name}")
            src_rel_root.append(new_rel)
            rel_map[old_id] = new_id

        new_elements = []
        for el in list(frag_body):
            if el.tag == f"{{{W_NS}}}sectPr":
                continue
            clone = copy.deepcopy(el)
            for node in clone.xpath(".//*[@r:embed]", namespaces=NS):
                old = node.get(f"{{{R_NS}}}embed")
                if old in rel_map:
                    node.set(f"{{{R_NS}}}embed", rel_map[old])
            # Reuse the source document's heading/caption style IDs.
            for ps in clone.xpath(".//w:pStyle", namespaces=NS):
                val = ps.get(f"{{{W_NS}}}val")
                if val == "Heading1":
                    ps.set(f"{{{W_NS}}}val", "1")
                elif val == "Heading2":
                    ps.set(f"{{{W_NS}}}val", "2")
                elif val == "Caption":
                    ps.set(f"{{{W_NS}}}val", "ImageCaption")
            new_elements.append(clone)

        for el in children[start:end]:
            src_body.remove(el)
        insert_at = start
        for el in new_elements:
            src_body.insert(insert_at, el)
            insert_at += 1

        src_tree.write(str(src_dir / "word/document.xml"), xml_declaration=True,
                       encoding="UTF-8", standalone=True)
        src_rels.write(str(rel_path), xml_declaration=True, encoding="UTF-8",
                       standalone=True)

        # Add TableGrid style if the source package does not already contain it.
        src_styles_path = src_dir / "word/styles.xml"
        frag_styles_path = frag_dir / "word/styles.xml"
        src_styles = etree.parse(str(src_styles_path), parser)
        frag_styles = etree.parse(str(frag_styles_path), parser)
        if not src_styles.xpath(".//w:style[@w:styleId='TableGrid']", namespaces=NS):
            table_grid = frag_styles.xpath(".//w:style[@w:styleId='TableGrid']", namespaces=NS)[0]
            src_styles.getroot().append(copy.deepcopy(table_grid))
            src_styles.write(str(src_styles_path), xml_declaration=True,
                             encoding="UTF-8", standalone=True)

        if output.exists():
            output.unlink()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as z:
            for p in sorted(src_dir.rglob("*")):
                if p.is_file():
                    z.write(p, p.relative_to(src_dir))


def main() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    configure_plotting()
    figures = {
        "b4": ASSET_DIR / "figure6_b4_negrec_scope.png",
        "d7d11": ASSET_DIR / "figure7_d7_d11_missing_arm.png",
        "d8": ASSET_DIR / "figure8_d8_threshold_decision.png",
    }
    make_b4_figure(figures["b4"])
    make_d7_d11_figure(figures["d7d11"])
    make_d8_figure(figures["d8"])
    with tempfile.TemporaryDirectory() as tmp:
        fragment = Path(tmp) / "part4_fragment.docx"
        build_fragment(fragment, figures)
        merge_fragment(SOURCE, fragment, OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
