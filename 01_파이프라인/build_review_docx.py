# -*- coding: utf-8 -*-
"""
'현 연구의 모델 개발 방향 및 데이터 수집에 관한 종합적 고찰' 재정리본.
원문(6개 절 순서) → 아이디어 6개 축으로 재구성. 어투는 원문의 학술적 평서문(-다) 유지.
그림은 원문에 삽입돼 있던 4개를 그대로 재사용(report_assets/orig_figs/).
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = "/Users/hanseoyun/Desktop/260830"
FIG = os.path.join(BASE, "report_assets", "orig_figs")
OUT = os.path.join(BASE, "종합적_고찰_정리본_20260830.docx")

KFONT = "맑은 고딕"
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x2C, 0x3E, 0x50)
GREY = RGBColor(0x7F, 0x8C, 0x8D)

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.0)
sec.left_margin = sec.right_margin = Cm(2.2)

st = doc.styles["Normal"]
st.font.name = KFONT
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), KFONT)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.32

for nm, sz in [("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)]:
    s = doc.styles[nm]
    s.font.name = KFONT
    s.element.rPr.rFonts.set(qn("w:eastAsia"), KFONT)
    s.font.size = Pt(sz)
    s.font.color.rgb = BLUE
    s.font.bold = True


def _run(p, text, bold=False, size=10.5, color=None, italic=False):
    r = p.add_run(text)
    r.font.name = KFONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), KFONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def h1(t):
    doc.add_paragraph(t, style="Heading 1").paragraph_format.space_before = Pt(18)


def h2(t):
    doc.add_paragraph(t, style="Heading 2").paragraph_format.space_before = Pt(12)


def h3(t):
    doc.add_paragraph(t, style="Heading 3").paragraph_format.space_before = Pt(8)


def para(*parts, align=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for x in parts:
        if isinstance(x, tuple):
            _run(p, x[0], **x[1])
        else:
            _run(p, x)
    return p


def bullet(text, lvl=0, bold_head=None):
    p = doc.add_paragraph(style="List Bullet" if lvl == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.6 + 0.5 * lvl)
    if bold_head:
        _run(p, bold_head, bold=True)
    _run(p, text)
    return p


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def table(headers, rows, widths=None, head_fill="2E74B5", zebra="F0F4F8", font=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, hh in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, hh, bold=True, size=font, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(hdr[i], head_fill)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, v in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            _run(p, str(v), size=font)
            if zebra and ri % 2 == 1:
                shade(cells[ci], zebra)
    if widths:
        t.autofit = False
        for r in t.rows:
            for c, w in zip(r.cells, widths):
                c.width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(fname, caption, width=13.5):
    doc.add_picture(os.path.join(FIG, fname), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para((caption, dict(size=9, color=GREY, italic=True)),
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)


def note(text, fill="EAF2F8"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    c.text = ""
    _run(c.paragraphs[0], text, size=9.5, color=DARK)
    shade(c, fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ══════════════════════════════════════════════════════════════
# 표지
# ══════════════════════════════════════════════════════════════
doc.add_paragraph().paragraph_format.space_after = Pt(30)
para(("혼합물(농약 제형) 안구·피부자극 평가에서", dict(bold=True, size=17, color=BLUE)),
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para(("GHS 농도한계값(CT) 접근법의 가산성 가정 한계와 대응 전략", dict(bold=True, size=17, color=BLUE)),
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
para(("아이디어 중심 정리본", dict(size=11, color=GREY)),
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)

para("농약 제형과 같은 혼합물의 안구·피부 자극 잠재력을 GHS 농도한계값(Concentration Threshold, "
     "CT) 접근법으로 예측할 때, 성분 간 상호작용(상승·길항작용)으로 단순 가산성 가정이 무너질 수 "
     "있다는 문제를 정리하고, 실험적 검증의 현실적 한계와 대응 방안, 그리고 이를 반영한 QSAR "
     "모델 개발 전략을 제시한다.")

h2("이 문서를 관통하는 아이디어 6개")
for i, t in enumerate([
    ("GHS CT 접근법은 \"성분들이 서로 영향을 주지 않는다\"는 가산성 가정 위에서 작동하는데, "
     "실제 혼합물에서는 이 가정이 깨질 수 있다."),
    ("이는 이론적 우려가 아니라, 실증 연구에서 절반 이상의 사례로 확인되는 문제다."),
    ("GHS 스스로도 일부 성분(강산·강염기·계면활성제 등)에서는 가산성이 통하지 않는다는 것을 "
     "인지하고 예외 조항을 두고 있다."),
    ("그러나 이 예외 조항은 이미 알려진 상호작용만 걸러내며, 성분 조합이 늘어날수록 모든 경우를 "
     "실험으로 확인하는 것은 현실적으로 불가능하다."),
    ("따라서 대응의 방향은 CT 접근법을 버리는 것이 아니라, 다른 근거(문헌·시험관내 자료·QSAR)와 "
     "함께 쓰는 다층적 필터링 구조로 CT의 역할을 재배치하는 것이다."),
    ("이 논리를 실제 모델 개발 계획으로 옮기면, CT를 거치지 않는 직접예측 트랙(Track A)과 "
     "CT를 기반으로 두고 오차만 보정하는 잔차보정 트랙(Track B)을 병행하는 전략이 나온다."),
], 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(5)
    _run(p, f"아이디어 {i}.  ", bold=True, color=BLUE)
    _run(p, t)

doc.add_page_break()

# ══════════════════════════════════════════════════════════════
# 아이디어 1
# ══════════════════════════════════════════════════════════════
h1("아이디어 1. 가산성 가정은 무엇이고, 왜 깨질 수 있는가")

para("농약 제형은 통상 1개 이상의 유효성분과 다수의 보조성분(계면활성제, 용매, 안정제 등)으로 "
     "구성된 혼합물이다. GHS CT 접근법은 각 성분이 자신의 농도·개별 역가에 비례적으로만 전체 "
     "제형의 자극성에 기여한다는 가산성(additivity) 가정 위에서 작동한다.")

para("그러나 성분 a가 단독으로는 안정(비자극)하고 성분 b가 단독으로는 독성(자극)을 나타내더라도, "
     "a와 b가 특정 농도비로 혼합되면 다음 세 가지 결과 중 하나가 나타날 수 있다.")

bullet("a+b의 자극성이 각 성분 기여도의 단순 합과 일치.", bold_head="① 상가적 반응 — 예측대로: ")
bullet("a+b의 실제 자극성이 예측보다 훨씬 큼. 개별로는 안전한 조합이 위험해지는 경우.",
       bold_head="② 상승작용(synergism): ")
bullet("a+b의 실제 자극성이 예측보다 작음. 한 성분이 다른 성분의 효과를 상쇄하는 경우.",
       bold_head="③ 길항작용(antagonism): ")

para("문제는 ②와 ③이 나타나는지, 나타난다면 그 정도가 얼마인지가 성분의 종류뿐 아니라 정확한 "
     "농도비에 의존한다는 점이다. 이 때문에 모든 조합·모든 배합비를 생체내 시험으로 사전에 전수 "
     "검증하는 것은 현실적으로 불가능하다.")

h2("개념 도구 — 등독성선(Isobologram)")

para("이 현상을 가장 직관적으로 보여주는 도구가 등독성선(isobologram)이다. 동일한 독성 효과를 "
     "발생시키는 성분 a, b의 모든 조합을 좌표평면에 표시하면, 가산성 가정 하에서는 두 성분의 "
     "단독 유효용량을 잇는 직선이 나타난다. 실제 관찰된 조합의 궤적이 이 직선보다 원점 쪽으로 "
     "휘어지면 상승작용, 바깥쪽으로 휘어지면 길항작용을 의미한다.")

figure("fig_isobologram.png",
       "그림 1. 등독성선 개념도 — 상가작용 예측선(검은 실선)과 실제 상승·길항작용 궤적의 차이.")

para("그림 1에서 보듯, GHS CT 접근법은 항상 검은 실선(상가작용 예측선)만을 계산하도록 설계돼 "
     "있다. 실제 제형의 반응이 붉은 점선(상승) 영역에 속한다면 CT 접근법은 위험을 과소평가하게 "
     "되고, 파란 점선(길항) 영역에 속한다면 반대로 과대평가(불필요한 동물시험 유발)로 이어진다.")

# ══════════════════════════════════════════════════════════════
# 아이디어 2
# ══════════════════════════════════════════════════════════════
doc.add_page_break()
h1("아이디어 2. 실제로 얼마나 자주 일어나는가 — 세 갈래 증거")

para("가산성 가정의 붕괴는 이론적 가능성에 그치지 않는다. 통계적 증거, 제도적 인정, 개별 사례 "
     "세 갈래에서 모두 확인된다.")

h2("증거 ① 통계 — 절반 이상이 비가산적")

para("국내 하천의 농약-의약품 2성분 혼합물 60개 지점을 분석한 연구에서는, 예측 대비 관찰 독성의 "
     "비율(Model Deviation Ratio)을 기준으로 상가작용 48% · 길항작용 40% · 상승작용 12%로 나타나, "
     "절반 이상의 사례에서 단순 가산성 예측이 실제와 어긋났다.")

figure("fig_60sites.png",
       "그림 2. 이성분 혼합물 60개 지점의 독성 상호작용 유형 분포 (예시 연구 결과).")

para("더 나아가 같은 연구에서는 저농도(EC10) 구간에서 상승작용을, 고농도(EC50) 구간에서는 오히려 "
     "길항작용을 보이는 역전 현상도 일부 조합에서 관찰됐다. 이는 상호작용의 방향과 크기가 농도 "
     "구간에 따라 달라질 수 있음을 의미하며, 단일 농도점에서의 검증만으로는 전체 배합비 스펙트럼의 "
     "안전성을 보장할 수 없다는 것을 시사한다.")

para("농약 제형 자체에 대해서도 유사한 우려가 보고된 바 있다. 급성 전신독성 영역에서는 제형의 "
     "독성이 성분 독성의 단순 합이 아니며, 성분 간 상호작용으로 GHS 가산성 예측보다 높거나 낮은 "
     "독성을 보일 수 있다는 점이 CT 접근법 개발자들 스스로에 의해 명시적으로 인정된 바 있다.")

h2("증거 ② 제도 — GHS 자신이 인정한 예외")

para("주목할 점은 GHS 제정 당시부터 특정 화학물질군에 대해서는 가산성 접근법이 통하지 않는다는 "
     "것이 이미 인지돼 있었다는 것이다. UN GHS 10차 개정판 및 이를 반영한 각국 규정(OSHA HCS, "
     "EU CLP 등)은 다음과 같은 예외를 명시한다.")

table(["구분", "해당 성분·상황", "대체 분류 기준"],
      [["강산·강염기", "pH ≤ 2 또는 pH ≥ 11.5", "농도가 아닌 pH를 우선 기준으로 사용"],
       ["무기염", "일부 무기염류", "개별 물질별 실측 자료 우선"],
       ["알데히드·페놀", "반응성이 높아 예측 곤란한 화학군", "GHS Table 3.2.4 / 3.3.4 등 별도 표 적용"],
       ["계면활성제", "피부·안구 장벽 투과를 촉진할 수 있는 성분", "가산성 배제, 실측 우선 권고"]],
      widths=[3.0, 5.8, 5.8], font=9)

para("즉 가산성 문제는 GHS 체계가 이미 알고 있는 한계이며, 그래서 위와 같은 예외 조항이 존재한다. "
     "다만 이 조항은 상승·길항 가능성이 알려진 특정 화학군에만 한정돼 있어, 문헌에 보고되지 않은 "
     "새로운 성분 조합이나 미지의 상호작용까지는 포괄하지 못한다는 것이 실무적 한계로 남는다.")

h2("증거 ③ 사례 — 농약 제형에서 보고된 구체적 상호작용")

table(["사례", "방향", "기전(가설)", "시사점"],
      [["계면활성제 SDS + C12TAB 혼합", "길항",
        "양이온성 계면활성제가 음이온성 계면활성제의 단량체 농도를 낮춰 피부장벽 교란력 감소",
        "계면활성제끼리도 서로 억제 가능 — 단순 합산은 과대예측 위험"],
       ["글리포세이트 + 보조제(POEA 등)", "상승",
        "보조제가 활성성분의 세포막 투과·흡수를 촉진, 산화스트레스·세포독성 증폭",
        "활성성분 단독 자료만으로는 제형의 실제 위해도를 과소평가할 위험"],
       ["피페로닐부톡사이드(공동상승제) + 살충제", "상승 (의도적)",
        "P450 대사효소 억제로 살충제 대사 지연, 효력·독성 모두 증대",
        "일부 제형은 상승작용을 의도적으로 설계 — 우연이 아닌 구조적 문제"],
       ["농약-의약품 이성분 혼합물 (하천수 사례)", "혼재 (48% / 40% / 12%)",
        "화학종별 작용기전 상이로 조합에 따라 결과 상이",
        "혼합물 유형의 일반화가 어려움 — 사례별 검증 필요성 시사"]],
      widths=[3.6, 2.0, 5.4, 3.8], font=8.5)

# ══════════════════════════════════════════════════════════════
# 아이디어 3
# ══════════════════════════════════════════════════════════════
doc.add_page_break()
h1("아이디어 3. 그런데 왜 실험으로 전부 확인할 수 없는가")

para("여기서 핵심은 상호작용이 존재한다는 사실 자체가 아니라, 이를 실험으로 전수 검증하는 것이 "
     "비현실적이라는 데 있다. 이유는 네 가지로 정리된다.")

bullet("성분이 n개이면 2성분 조합만 해도 nC2가지이며, 실제 제형은 성분이 10개 이상인 경우가 흔해 "
       "조합 수가 기하급수적으로 증가한다.", bold_head="성분 수 증가에 따른 조합 폭증: ")
bullet("상호작용의 방향(상승·길항)과 크기가 배합비에 따라 달라지므로, 하나의 배합비에서 안전을 "
       "확인해도 다른 배합비에서는 다른 결과가 나올 수 있다.", bold_head="농도 의존성: ")
bullet("저농도에서 상승, 고농도에서 길항으로 방향이 바뀌는 사례가 보고돼, 단일 시험 조건만으로는 "
       "전체 노출 스펙트럼을 대표할 수 없다.", bold_head="농도 구간별 역전 현상: ")
bullet("계면활성제·대사억제 등 알려진 상호작용은 예외 조항으로 걸러낼 수 있지만, 아직 보고되지 "
       "않은 조합·기전은 사전에 예측하기 어렵다.", bold_head="기전 미상의 상호작용: ")

note("정리하면, \"검증이 불가능하다\"는 사실 자체가 대응 전략의 출발점이다. 전수 실험이라는 이상적 "
     "해법이 막혀 있는 상태에서, 가진 근거를 최대한 조합하는 현실적 전략이 다음 아이디어의 내용.")

# ══════════════════════════════════════════════════════════════
# 아이디어 4
# ══════════════════════════════════════════════════════════════
doc.add_page_break()
h1("아이디어 4. 그래서 어떻게 대응하는가 — 폐기가 아니라 재배치")

para("현재로서는 가산성 가정의 근본적 한계를 완전히 해소하는 단일 해법은 없다. 다만 다음과 같은 "
     "다층적 보완 전략이 실무·학계에서 제안되고 있다.")

table(["전략", "핵심 내용", "한계"],
      [["가중치 증거 (Weight of Evidence)",
        "CT 계산 결과를 단독 근거로 쓰지 않고 시험관내 시험(RhE, BCOP, ICE 등) 결과와 교차 검증",
        "시험관내 자료가 없는 신규 제형에는 적용 곤란"],
       ["알려진 상호작용 성분 배제 규칙",
        "계면활성제·산·염기 등 GHS가 이미 지정한 비가산성 성분 포함 시 CT 계산에서 제외하고 실측 우선",
        "미지의 상호작용은 걸러내지 못함"],
       ["물리화학적 특성 추가 반영",
        "계면활성제 함량(%), 용매 극성 등 침투능에 영향을 주는 변수를 별도로 고려",
        "정량적 보정계수에 대한 합의 부족"],
       ["QSAR/기전기반 모델링",
        "농도가산모형(CA)·독립작용모형(IA) 양쪽에 기반한 다중회귀·기계학습 모델로 이성분 조합의 상호작용 예측",
        "복잡한 다성분 제형에는 적용 사례가 아직 제한적, 적용영역(AD) 정의 필요"],
       ["PBTK/AOP 기반 통합평가(IATA)",
        "생리학적기반 동태모델(PBTK), 유해결과경로(AOP)를 결합해 단일물질 자료를 넘어서는 기전기반 통합 평가",
        "자료·모델 구축에 상당한 자원 소요, 표준화 미흡"]],
      widths=[3.6, 6.6, 4.6], font=8.5)

h2("이를 하나의 흐름으로 통합하면")

para("위 대안들을 실무에 적용 가능한 형태로 통합하면, 아래와 같은 단계적 의사결정 흐름을 제안할 "
     "수 있다.")

figure("fig_decision_flow.png",
       "그림 3. 혼합물(농약 제형) 안구·피부자극 평가를 위한 단계적 의사결정 흐름도 (제안).",
       width=11.5)

para("이 흐름도의 핵심은, CT 접근법을 전면 신뢰하거나 전면 불신하는 이분법이 아니라, "
     "① GHS가 이미 인지한 예외 성분을 먼저 걸러내고, ② 문헌 기반으로 알려진 상호작용 성분을 "
     "추가로 스크리닝하며, ③ 가능한 경우 시험관내 자료와 교차검증하고, ④ 근거가 일치할 때만 "
     "동물시험 면제를 고려하는 다층적 필터링 구조에 있다.")

h2("결론 — CT는 버릴 도구가 아니라 재배치할 도구")

bullet("가산성 가정은 화학물질이 서로 독립적으로 작용한다는 전제에 기반하지만, 실제 제형에서는 "
       "대사적 상호작용, 흡수·투과 촉진, 물리화학적 길항(미셀형성 등)으로 이 전제가 깨질 수 있다.")
bullet("이는 이론적 우려가 아니라 여러 실증 연구(60개 혼합물 지점 중 52%가 비가산적)에서 확인된 "
       "문제이며, GHS 규정 자체도 일부 성분군에 대해 이를 인정하고 예외 조항을 두고 있다.")
bullet("그러나 예외 조항은 알려진 상호작용에만 대응할 뿐, 성분 수·배합비 조합이 기하급수적으로 "
       "늘어나는 실제 제형에서 모든 상호작용을 사전 실험으로 검증하는 것은 현실적으로 불가능하다.")
bullet("따라서 CT 접근법을 완전 대체 도구가 아닌 1차 스크리닝 + 가중치 증거 구성요소로 자리매김 "
       "시키고, 알려진 위험 성분 배제 규칙, 시험관내 교차검증, 기전기반 모델링(QSAR/PBTK/AOP)을 "
       "병행하는 다층적 접근이 현재 가장 현실적인 대응 전략이다.")

para("정리하면, GHS CT 접근법의 근본적 가정(가산성)에 대한 문제 제기는 정당한 과학적 비판이며, "
     "이는 현재 국제적으로도 미해결 과제로 남아 있다. 다만 이것이 CT 접근법 자체의 폐기를 의미하지 "
     "않으며, 오히려 어떤 조건에서 CT 접근법을 신뢰할 수 있는가를 정교화하는 방향(가중치 증거, "
     "예외 성분 배제, 기전기반 모델 통합)으로 관련 연구가 진행되고 있다는 점이 중요하다.")

# ══════════════════════════════════════════════════════════════
# 아이디어 5
# ══════════════════════════════════════════════════════════════
doc.add_page_break()
h1("아이디어 5. 이 논리를 모델 개발로 옮기면 — Track A / Track B 병행")

para("여기서부터는 앞선 논의(가산성 가정의 한계)를 실제 QSAR 모델 개발 계획으로 구체화한 것이다. "
     "현재 데이터 수집이 진행 중인 단계임을 고려해, (A) 데이터를 최대한 수집해 관찰값을 직접 "
     "예측하는 트랙과 (B) GHS CT 예측값을 baseline으로 두고 그 편차만을 학습하는 하이브리드 트랙을 "
     "병행 개발하는 전략을 제안한다. 두 트랙은 동일한 원자료(성분 SMILES·농도·pH·관찰 GHS 라벨)를 "
     "공유하므로 수집 파이프라인을 이원화할 필요는 없다.")

figure("fig_trackab.png", "그림 4. QSAR 모델 개발을 위한 2개 트랙(직접 예측 / CT 보완 하이브리드) 병행 전략도.")

h2("Track A — 직접 예측 (데이터 최대 수집)")

para("Track A는 CT 예측값을 참조하지 않고, 성분 구조·농도·pH 등 원 특징으로부터 관찰된 GHS 분류를 "
     "곧바로 학습한다. 이 트랙의 성패는 전적으로 데이터 양과 클래스 균형에 달려 있으므로, 수집 "
     "작업을 다음 두 방향으로 최대화할 필요가 있다.")

bullet("ECHA CHEM(REACH 등록자료), US EPA ECOTOX Knowledgebase, PubChem BioAssay 등에서 확보 가능한 "
       "이성분·다성분 혼합물 자극/독성 자료를 추가 수집해 표본 수를 늘린다.",
       bold_head="공개 데이터베이스 확충: ")
bullet("혼합물 라벨 데이터는 절대적으로 희소하므로, 먼저 단일 성분 단위의 대규모 독성·물성 "
       "데이터(logP, 피부투과성 등)로 사전학습한 인코더를 혼합물 문제에 미세조정하는 방식이 최근 "
       "환경독성 분야에서도 효과적으로 보고되고 있다. 기전적으로 관련된 대규모 baseline 성질(예: "
       "소수성)로 사전학습한 그래프신경망을 소규모 실험 데이터셋에 전이했을 때, 훨씬 큰 비관련 "
       "데이터셋으로 사전학습한 경우와 동등한 성능을 보였다는 연구가 있다.",
       bold_head="단일물질 QSAR로부터의 전이학습: ")

para("다만 Track A는 실제 제형의 70~85%가 GHS 미분류에 쏠려 있는 근본적인 클래스 불균형 문제를 "
     "안고 있으므로, 다음과 같은 불균형 처리 기법을 반드시 병행해야 한다.")

table(["기법", "방식", "독성 QSAR 분야 근거"],
      [["class weighting", "손실함수에서 소수 클래스(상승·길항 사례)에 더 큰 가중치 부여",
        "추가 데이터 생성 없이 즉시 적용 가능한 기본 옵션"],
       ["SMOTE (oversampling)", "소수 클래스 샘플 사이를 보간해 합성 샘플 생성",
        "독성 유전독성 데이터셋에서 분자 지문(MACCS)과 결합한 SMOTE가 가장 높은 F1 점수를 기록"],
       ["SMOTEENN", "SMOTE 오버샘플링 후 이웃 기반 정제로 잘못 라벨된 경계 샘플 제거",
        "Tox21과 같은 고도 불균형 독성 데이터셋에서 성능 개선이 보고됨"],
       ["undersampling", "다수 클래스(NC) 샘플 수를 줄여 균형 확보",
        "정보 손실 위험이 있으나 소수 클래스 특징이 뚜렷할 때 유효"],
       ["결정 임계값 튜닝", "분류 확률의 판정 경계(threshold)를 데이터 분포에 맞게 재조정",
        "재샘플링 없이 소수 클래스 재현율을 높이는 보완적 기법으로 제안됨"]],
      widths=[3.0, 5.4, 6.4], font=8.5)

h2("Track B — GHS CT 보완 하이브리드 모델")

para("Track B는 CT 접근법을 버리지 않고 그 위에 보정 모델을 얹는 구조다. 이는 물리·화학 등 타 "
     "분야에서 이미 검증된 잔차학습(residual learning) 프레임워크와 원리적으로 동일하다. "
     "예컨대 원자핵 붕괴반감기 예측에서 거시적 물리모델(ELDM)을 baseline으로 두고 머신러닝이 그 "
     "예측오차만을 보정하는 접근이나, 레이저피닝 잔류응력 예측에서 저정밀 해석모델을 신경망으로 "
     "보정하는 접근에서, 이러한 baseline+보정 구조가 순수 데이터기반 모델보다 적은 표본으로도 "
     "우수한 성능을 냄이 확인된 바 있다.")

para("독성학 분야에서도 이미 유사한 하이브리드 설계가 시도되고 있다. 조류 대상 아졸계 살균제 "
     "혼합독성 연구에서는 농도가산(CA)·독립작용(IA) 예측값을 분자기술자와 함께 입력변수로 사용한 "
     "모델이 가장 높은 예측력을 보였으며, 항생제-살균제 혼합물 연구에서도 QSAR 기술자에 "
     "유사도기반(RASAR) 및 잔차기반(ARKA) 기술자를 결합한 하이브리드 모델이 단독 QSAR보다 향상된 "
     "품질을 보였다. 병원성 화학물질 혼합물 독성평가에서도 병태생리 모델과 AI를 결합한 하이브리드 "
     "프레임워크가 단독 AI 모델의 한계를 보완한다고 보고됐다.")

para("Track B의 구체적 설계는 다음과 같다.")

bullet("관찰 GHS 분류(또는 관찰 자극점수) − CT 예측 분류(또는 CT 계산 점수). 연속값이라면 "
       "log(MDR)과 유사한 편차 지표로 정의 가능.", bold_head="타겟 변수: ")
bullet("순수 농도가중 항은 이미 CT가 포착하므로 제외하고, 상호작용 항(계면활성제 간 전하조합, "
       "logP 차이, 대사효소 억제 곱 등 — 아이디어 4의 QSAR/기전기반 특징)만을 집중 사용한다.",
       bold_head="입력 특징: ")
bullet("표본이 적은 초기 단계에서는 트리 기반 모델(RF/XGBoost)이 잔차 학습에 안정적이며, 표본이 "
       "축적되면 신경망 기반 보정으로 확장 가능하다.", bold_head="모델: ")

note("Track B의 가장 큰 장점은 데이터가 적은 현재 상태에서도 즉시 실행 가능하다는 점이다 — "
     "CT 예측값은 추가 실험 없이 계산만으로 확보되므로, 관찰-예측 쌍만 있으면 바로 잔차 모델 학습을 "
     "시작할 수 있다.")

h2("두 트랙 비교 및 선택 기준")

table(["항목", "Track A (직접 예측)", "Track B (CT 보완 하이브리드)"],
      [["필요 데이터량", "많음 — 전체 분포를 학습해야 하므로 데이터 최대화 필수",
        "상대적으로 적음 — 편차만 학습하므로 신호가 좁고 명확"],
       ["현 단계 실행 가능성", "데이터 수집이 완료된 이후 본격 성능 확보 가능",
        "현재 데이터로도 즉시 시작 가능 (CT는 계산만으로 확보)"],
       ["해석 가능성", "상가항과 상호작용항이 뒤섞여 원인 분리가 어려움",
        "편차의 원인이 곧 상호작용이므로 SHAP 등으로 기전 해석 용이"],
       ["불균형 문제", "심각 (NC 70~85% 편중) — SMOTE 등 필수",
        "상대적으로 완화 (상가적 사례가 편차 0 근방에 몰려도 회귀로 처리 가능)"],
       ["기존 CT 자산 활용", "활용 안 함 — 처음부터 다시 학습",
        "최대한 활용 — 검증된 baseline 위에 증분 개선"],
       ["리스크", "데이터 부족 시 CT를 단순 재현하는 모델로 수렴할 위험",
        "CT 자체의 구조적 오류(예: 산·염기 오분류)는 그대로 물려받음"]],
      widths=[2.6, 6.0, 6.2], font=8.5)

para("결론적으로 두 트랙은 상호 배타적이지 않다. 현재 데이터가 부족한 시점에서는 Track B로 즉시 "
     "개발에 착수해 조기 성과(baseline 대비 개선폭)를 확보하고, 데이터 수집이 누적됨에 따라 "
     "Track A의 성능을 함께 끌어올려, 최종적으로 두 트랙의 예측을 앙상블(가중평균, 스태킹 등)하는 "
     "것이 가장 현실적인 개발 로드맵이다.")

# ══════════════════════════════════════════════════════════════
# 아이디어 6
# ══════════════════════════════════════════════════════════════
doc.add_page_break()
h1("아이디어 6. 실행 로드맵")

table(["단계", "내용"],
      [["Phase 1 (즉시 착수, Track B 우선)",
        "현재 수집된 성분·농도·라벨 자료에 GHS CT 계산값을 추가 산출 → 편차 타겟 정의 → "
        "RF/XGBoost 잔차모델 1차 구축"],
       ["Phase 2 (병행)",
        "공개 DB(ECHA, ECOTOX, PubChem) 추가 수집 + 단일물질 QSAR 전이학습으로 Track A용 데이터 "
        "확충, class weighting·SMOTE 적용해 1차 직접예측 모델 구축"],
       ["Phase 3",
        "두 트랙 성능을 동일한 홀드아웃 세트(및 SDS+C12TAB형 알려진 상호작용 사례)에서 비교 검증"],
       ["Phase 4",
        "두 트랙 예측을 앙상블하거나, Track B의 해석성과 Track A의 포괄성을 결합한 최종 모델로 통합"]],
      widths=[4.4, 10.4], font=9)

para(("현재 진행 상황", dict(bold=True)),
     " — Phase 1(CT baseline 산출·Track B 잔차 정의)과 Phase 2(라벨 수집 확충·불균형 처리)가 "
     "동시에 진행 중이며, 구체적인 데이터 현황과 성능 수치는 별도 문서(데이터 현황·진척 보고)에 "
     "정리돼 있다.")

h1("참고 문헌 및 자료 출처")

para("본 정리본은 다음 유형의 자료를 종합적으로 검토하여 작성됐다: (1) UN GHS 10차 개정판 및 "
     "각국 이행 지침(OSHA HCS, EU CLP)의 혼합물 분류 규정, (2) 환경독성학 분야의 혼합물 상호작용 "
     "실증 연구(Model Deviation Ratio 기반 상승·길항 판별 연구), (3) 농약 제형의 보조제·활성성분 "
     "상호작용에 관한 독성학 문헌, (4) GHS 가산성 공식의 예측력에 관한 회고적 검증 연구(급성 "
     "전신독성 영역), (5) 신규접근법(NAM), QSAR, PBTK, IATA 등 대안적 혼합물 평가 방법론에 관한 "
     "리뷰 논문, (6) 클래스 불균형 처리 기법(SMOTE, SMOTEENN, 임계값 튜닝)에 관한 화학정보학·독성 "
     "QSAR 문헌, (7) 잔차·보정 기반 하이브리드 모델링(물리기반 baseline + 머신러닝 보정) 및 소규모 "
     "데이터 전이학습에 관한 연구. 구체적 수치·사례는 해당 분야의 공개된 학술 문헌에서 확인된 "
     "값을 인용·재구성한 것이다.")

doc.save(OUT)
print("saved →", OUT)
print("size  =", round(os.path.getsize(OUT) / 1024, 1), "KB")
