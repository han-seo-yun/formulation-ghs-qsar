# -*- coding: utf-8 -*-
"""
이번 주(8/24~) 데이터 추가 보완 작업 지시 파일.
시트는 '이번 작업 항목' 기준으로 분류(담당자 기준 아님) — 사용자 지시 반영.
원본 데이터는 건드리지 않고 새 워크북으로 산출.
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "/Users/hanseoyun/Desktop/260830/이번주_작업지시_20260824.xlsx"

NAVY = "1A5276"
GREY = "7F8C8D"
LIGHT = "EAF2F8"
WHITE = "FFFFFF"

wb = Workbook()
wb.remove(wb.active)

HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(color=WHITE, bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
SUB_FONT = Font(bold=True, size=11, color=NAVY)
BODY_FONT = Font(size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row=1, ncols=None):
    ncols = ncols or ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 28


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_df(ws, df, start_row=1, widths=None):
    for j, col in enumerate(df.columns, start=1):
        ws.cell(start_row, j, col)
    style_header(ws, start_row, len(df.columns))
    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            cell = ws.cell(i, j, "" if pd.isna(v) else v)
            cell.border = BORDER
            cell.alignment = WRAP if j in (3, 4) else Alignment(vertical="top")
    if widths:
        autosize(ws, widths)
    ws.freeze_panes = ws.cell(start_row + 1, 1).coordinate


# ══════════════════════════════════════════════════════════════
# 00_작업안내
# ══════════════════════════════════════════════════════════════
ws = wb.create_sheet("00_작업안내")
r = 1
ws.cell(r, 1, "이번 주 데이터 추가 보완 작업 — 검증 결과 및 분담").font = TITLE_FONT
r += 2
ws.cell(r, 1, "1. 이번 주 검증 결과").font = SUB_FONT
r += 1
summary_rows = [
    ("성분 시트 — 팀 판정 유지 + 규칙 교차검증",
     "기존 팀 판정은 그대로 유지. 규칙 판정과 대조해 판정이 다른 행만 별도 표시."),
    ("　· 팀 판정과 규칙 판정 일치", "1,624행"),
    ("　· 재확인 필요(판정 불일치)", "481행 → 시트 [01_성분판정_재확인]"),
    ("　· 제외 처리 대상", "218행 → 시트 [01_성분판정_재확인] (구분 열로 함께 표시)"),
    ("　· 아직 미검증(규칙 판정만 임시 적용)", "3,182행"),
    ("신규 확보 가능 (구조화 반영 시)",
     "CAS 약 500행, 함량 809행 추가 확보 가능. 기존 CAS-GHS 대응표와 겹쳐 성분 GHS 102행 즉시 복구 가능"),
    ("결과 파일", "dataset_배정_검증완료_20260824.xlsx (원본 dataset_배정_검증미완.xlsx 은 미수정)"),
]
for label, val in summary_rows:
    ws.cell(r, 1, label).font = Font(bold=not label.startswith("　"), size=10)
    ws.cell(r, 2, val).font = BODY_FONT
    ws.cell(r, 2).alignment = WRAP
    r += 1
r += 1

ws.cell(r, 1, "2. 이번 주 우선순위").font = SUB_FONT
r += 1
ws.cell(r, 1, "① pH 수집  →  ② SDS §11 라벨 회수  →  ③ 성분 GHS 및 판정 정리").font = Font(bold=True, size=10)
r += 1
ws.cell(r, 1, "P2/P4 백로그는 웹 검색이 필요한 데 비해 회수율이 낮아 이번 주는 우선순위를 낮춤.").font = BODY_FONT
r += 2

ws.cell(r, 1, "3. 담당별 작업").font = SUB_FONT
r += 1
tasks = [
    ("한서윤 (총괄)", "전체 파이프라인 재빌드 + 모델 학습 진행",
     "성분판정 불일치 481건 재확인, 제외 218건 처리기준 정리, 라벨충돌 222건 판정. "
     "대부분 규칙으로 1차 필터링 후 최종 확인하는 방식.",
     "01_성분판정_재확인, 02_라벨충돌_판정"),
    ("이서윤", "SDS Section 9 pH 수집 (액상 SL 코드 우선)",
     "약 280건. SDS에서 pH 또는 pH value 확인 후 입력.",
     "03_pH수집"),
    ("김보경", "SDS Section 9 pH 수집 (액상 SC/EC 코드 우선)",
     "약 274건. SDS에서 pH 또는 pH value 확인 후 입력.",
     "03_pH수집"),
    ("정채윤", "성분 GHS 분류 조사 (PubChem PUG-View / ECHA C&L)",
     "고유 CAS 331종 전담(중복 조사 방지를 위해 한 분이 전량 담당).",
     "04_성분GHS조사"),
    ("최소빈", "SDS Section 11 원문 재판독 (신규 웹검색 아님)",
     "총 797행 — 눈 268 · 피부 269 · 감작 260. 이미 확보된 원문에서 놓친 라벨 판독.",
     "05_SDS11_라벨판독"),
]
tdf = pd.DataFrame(tasks, columns=["담당", "작업", "세부 내용", "해당 시트"])
th = r
for j, col in enumerate(tdf.columns, start=1):
    ws.cell(th, j, col)
style_header(ws, th, 4)
for i, row in enumerate(tdf.itertuples(index=False), start=th + 1):
    for j, v in enumerate(row, start=1):
        c = ws.cell(i, j, v)
        c.border = BORDER
        c.alignment = WRAP
    ws.row_dimensions[i].height = 46
r = th + len(tdf) + 2

ws.cell(r, 1, "4. 공통 기록 규칙").font = SUB_FONT
r += 1
rules = [
    "값을 새로 채운 경우 근거 URL을 반드시 함께 남긴다.",
    "확인이 안 되는 값은 추정해서 채우지 않고 비워둔다.",
    "메모에 확인이 어려운 이유만 남긴다 — 태그: [SDS없음] [접근차단] [성분불일치] [제품값아님]",
    "전량 회수가 목표가 아니다 — 최대한 확인 후, 안 되는 건 사유만 정확히 남기면 충분하다.",
    "기준이 애매한 행은 임의로 판단하지 않고 '판단보류' 로 표시만 한다 (각 시트 메모/판정 열 참고).",
]
for x in rules:
    ws.cell(r, 1, "• " + x).font = BODY_FONT
    r += 1
r += 1
ws.cell(r, 1, "이번 라운드 반영이 끝나면 데이터 수집 단계의 큰 병목은 대부분 정리됩니다. "
              "이후 전체 파이프라인 재빌드와 모델 성능 확인은 한서윤이 진행합니다.").font = Font(italic=True, size=10, color=GREY)
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 46
ws.column_dimensions["D"].width = 22
ws.sheet_view.showGridLines = False

# ══════════════════════════════════════════════════════════════
# 01_성분판정_재확인 (한서윤)
# ══════════════════════════════════════════════════════════════
import openpyxl as oxl
src = oxl.load_workbook("/Users/hanseoyun/Desktop/260830/dataset_배정_검증완료_20260824.xlsx", data_only=True)
ing_ws = src["성분"]
hdr = [c.value for c in ing_ws[1]]
C = {h: i for i, h in enumerate(hdr) if h}

rows_conf, rows_excl = [], []
for rr in range(2, ing_ws.max_row + 1):
    method = ing_ws.cell(rr, C["검증방식"] + 1).value
    verdict = ing_ws.cell(rr, C["검증여부"] + 1).value
    rec = [
        ing_ws.cell(rr, C["Formulation_ID"] + 1).value,
        ing_ws.cell(rr, C["product_name"] + 1).value,
        ing_ws.cell(rr, C["ingredient_name"] + 1).value,
        ing_ws.cell(rr, C["추출정보"] + 1).value,
        verdict,
        ing_ws.cell(rr, C["검증근거"] + 1).value,
        "", "",  # 최종판정, 메모
    ]
    if method == "human(규칙불일치)":
        rows_conf.append(rec)
    elif verdict == "제외":
        rows_excl.append(rec)

cols = ["Formulation_ID", "product_name", "ingredient_name", "추출정보(현재값)",
        "구분", "검증근거(팀판정 vs 규칙판정)", "최종판정(한서윤 기입)", "메모"]
df1 = pd.DataFrame(rows_conf, columns=cols[:6] + cols[6:])
df1["구분"] = "판정불일치"
df2 = pd.DataFrame(rows_excl, columns=cols[:6] + cols[6:])
df2["구분"] = "제외후보"
df_all = pd.concat([df1, df2], ignore_index=True)[cols]

ws = wb.create_sheet("01_성분판정_재확인")
ws.cell(1, 1, "성분 시트 — 판정 불일치 481건 + 제외 후보 218건 (담당: 한서윤)").font = TITLE_FONT
write_df(ws, df_all, start_row=3, widths=[16, 26, 24, 42, 12, 46, 16, 22])
dv = DataValidation(type="list", formula1='"확정,판단보류"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"G4:G{3+len(df_all)}")

# ══════════════════════════════════════════════════════════════
# 02_라벨충돌_판정 (한서윤)
# ══════════════════════════════════════════════════════════════
conflict = pd.read_csv("/tmp/wk_conflict.csv")
src_col = [c for c in conflict.columns if c.endswith("_src")][0]
nsrc_col = [c for c in conflict.columns if c.endswith("_n_src")][0]
conflict = conflict.rename(columns={src_col: "라벨출처", nsrc_col: "출처수"})
conflict = conflict[["Formulation_ID","product_name","엔드포인트","현재라벨","라벨출처","출처수","할당","최종판정","메모"]]

ws = wb.create_sheet("02_라벨충돌_판정")
ws.cell(1, 1, "라벨 충돌 222건 — 눈 124 · 피부 88 · 감작 10 (담당: 한서윤)").font = TITLE_FONT
write_df(ws, conflict, start_row=3, widths=[16, 26, 10, 12, 16, 8, 10, 16, 22])
dv2 = DataValidation(type="list", formula1='"확정,판단보류"', allow_blank=True)
ws.add_data_validation(dv2)
dv2.add(f"H4:H{3+len(conflict)}")

# ══════════════════════════════════════════════════════════════
# 03_pH수집 (이서윤/김보경)
# ══════════════════════════════════════════════════════════════
ph = pd.read_csv("/tmp/wk_ph.csv")
ph = ph[["Formulation_ID","product_name","form_code_best","할당","pH_확보값","근거URL","메모"]]
ph.columns = ["Formulation_ID","product_name","제형코드","담당","pH_확보값(SDS Section 9)","근거URL","메모"]
ws = wb.create_sheet("03_pH수집")
ws.cell(1, 1, "SDS Section 9 pH 수집 — 554건 (이서윤 280 · 김보경 274)").font = TITLE_FONT
write_df(ws, ph, start_row=3, widths=[16, 30, 10, 8, 20, 40, 22])

# ══════════════════════════════════════════════════════════════
# 04_성분GHS조사 (정채윤)
# ══════════════════════════════════════════════════════════════
ghs = pd.read_csv("/tmp/wk_ghs.csv")
ghs = ghs[["CAS","대표성분명","관련행수","GHS_눈","GHS_피부","GHS_감작","근거URL","메모"]]
ws = wb.create_sheet("04_성분GHS조사")
ws.cell(1, 1, "성분 GHS 분류 조사 — 고유 CAS 331종 (담당: 정채윤, 전담)").font = TITLE_FONT
write_df(ws, ghs, start_row=3, widths=[16, 30, 10, 12, 12, 12, 40, 22])

# ══════════════════════════════════════════════════════════════
# 05_SDS11_라벨판독 (최소빈)
# ══════════════════════════════════════════════════════════════
s11 = pd.read_csv("/tmp/wk_s11.csv")
s11 = s11[["Formulation_ID","product_name","엔드포인트","원문(§11)","판독구분","메모"]]
ws = wb.create_sheet("05_SDS11_라벨판독")
ws.cell(1, 1, "SDS Section 11 원문 재판독 — 797건 (담당: 최소빈, 웹검색 불필요)").font = TITLE_FONT
write_df(ws, s11, start_row=3, widths=[16, 30, 10, 60, 14, 22])
dv3 = DataValidation(type="list", formula1='"구분확정,정보없음,판단보류"', allow_blank=True)
ws.add_data_validation(dv3)
dv3.add(f"E4:E{3+len(s11)}")

wb.save(OUT)
print("saved →", OUT)