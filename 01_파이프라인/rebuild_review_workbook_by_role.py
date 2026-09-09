#!/usr/bin/env python3
"""수동검토·최종판단 작업파일을 담당별 시트로 만든다.

구성
----
0_먼저읽기        담당별 시트 안내, 작업 순서, 입력 규칙
참조자료          R1~R18. 각 작업 행의 `참조` 열이 이 ID 를 가리킨다
규제1~3           관할 규정 확정 6건 / 피부 범주 정정 2건 / H코드 판정표
데이터1~2         원문 판독 196문서 / 개별 원문 확인 4건
모델1             코드 작업 4건
총책임자_결정      정책·방법론 결정 9건
참고_*            근거 자료. 채울 칸 없음

모든 작업 행에 `참조` 열이 있고, 참조자료 시트에서 경로·조항까지 찾을 수 있다.
읽기 전용 입력. 산출은 작업파일 1개.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
REV = ROOT / "04_모델산출물" / "v6_수동검토"
WB_IN = REV / "수동검토_최종판단_작업파일_이전판_260904백업.xlsx"
WB_OUT = REV / "dataset_배정_260904.xlsx"  # 구 파일명 수동검토_최종판단_작업파일.xlsx, 드라이브 업로드 후 이 이름으로 남음
DOCBASE = ROOT / "formulation_audit_team_share_highlighted_only_20260630"

H = {"규제": PatternFill("solid", fgColor="1F4E79"),
     "데이터": PatternFill("solid", fgColor="2E6B34"),
     "모델": PatternFill("solid", fgColor="7A4A00"),
     "총괄": PatternFill("solid", fgColor="7B1D1D"),
     "참고": PatternFill("solid", fgColor="595959")}
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_LOCK = PatternFill("solid", fgColor="F2F2F2")
WHITE = Font(color="FFFFFF", bold=True, size=10)
THIN = Border(*[Side("thin", color="BFBFBF")] * 4)


def _clean(v):
    """물려받은 시트 텍스트에 남아 있는 마크다운 표기를 벗긴다."""
    if not isinstance(v, str):
        return v
    return v.replace("**", "").replace("`", "").strip()


def sheet(wb, title, headers, rows, fill, fill_cols=(), widths=None,
          note=None, dv=None, wrap_cols=()):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        c = ws.cell(1, 1, note)
        c.font = Font(bold=True, size=10, color="1F4E79")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(start_row=1, start_column=1, end_row=1,
                       end_column=max(len(headers), 6))
        ws.row_dimensions[1].height = 32
        r0 = 3
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.fill, c.font = fill, WHITE
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
    ws.row_dimensions[r0].height = 30
    idx = {h: j + 1 for j, h in enumerate(headers)}
    for i, row in enumerate(rows, r0 + 1):
        for j, h in enumerate(headers, 1):
            v = row.get(h, "")
            c = ws.cell(i, j, "" if v is None or (isinstance(v, float)
                                                  and pd.isna(v)) else _clean(v))
            c.border = THIN
            c.alignment = Alignment(wrap_text=h in wrap_cols, vertical="top")
            c.fill = FILL_IN if h in fill_cols else FILL_LOCK
    for h, w in (widths or {}).items():
        if h in idx:
            ws.column_dimensions[get_column_letter(idx[h])].width = w
    for h in headers:
        if h not in (widths or {}):
            ws.column_dimensions[get_column_letter(idx[h])].width = 14
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    if rows:
        ws.auto_filter.ref = (f"A{r0}:"
                              f"{get_column_letter(len(headers))}{r0+len(rows)}")
    for cols, opts in (dv or []):
        v = DataValidation(type="list", formula1='"' + ",".join(opts) + '"',
                           allow_blank=True, showDropDown=False)
        ws.add_data_validation(v)
        for h in cols:
            if h in idx:
                L = get_column_letter(idx[h])
                v.add(f"{L}{r0+1}:{L}{r0+len(rows)}")
    return ws


# ========================================================== 원자료 회수
xl = pd.ExcelFile(WB_IN)
A, B, C, D, E, F = (xl.parse(s) for s in
                    ("A_관할규정_확정", "B_행단위_판별불가", "C_negrec_원문재수집",
                     "D_기타_미완항목", "E_L1정정_82셀", "F_부분확인_레지스트리"))
for s, d, n in (("A", A, 6), ("B", B, 4), ("C", C, 448), ("D", D, 10),
                ("E", E, 82), ("F", F, 34)):
    assert len(d) == n, f"{s} 행수 {len(d)} != {n}"
_filled = [c for c in ("판정결과(양성/음성/판별불가)", "H코드", "원문근거문장",
                       "확인자", "확인일") if int(C[c].notna().sum()) > 0]
assert not _filled, f"수기 입력이 있다 — 중단: {_filled}"

# ========================================================== 참조자료 레지스트리
# 경로는 전부 실존을 확인한 것만 적는다. 외부 규정은 발행처·조항까지만 적고
# 링크를 만들어 쓰지 않는다.
REFS = [
    ("R1", "이 파일 · 규제3_H코드_판정표", "이 워크북 내 시트",
     "H코드 → GHS 범주 → 관할별 양성/음성 대응", "데이터"),
    ("R2", "고용노동부 고시「화학물질의 분류·표시 및 물질안전보건자료에 관한 기준」",
     "국가법령정보센터에서 고시명 검색 → 별표 1 (유해성·위험성 분류 기준)",
     "눈 자극성 구분 2A/2B 세분 여부, 피부 자극성 구분 3 채택 여부", "규제"),
    ("R3", "OSHA HCS 29 CFR 1910.1200 Appendix A",
     "A.2 피부 부식성/자극성 · A.3 심한 눈 손상/눈 자극성 · A.4 호흡기·피부 과민성",
     "각 표에 Category 2B / Category 3 행이 있는지", "규제"),
    ("R4", "UN GHS 개정 10판 3.2 / 3.3", "3.2 피부 부식성·자극성 · 3.3 눈 손상·자극성",
     "선택범주(Cat 3, Cat 2B) 의 원 정의", "규제"),
    ("R5", "EU CLP 규정 (EC) No 1272/2008", "Annex I 3.2·3.3 / Annex VI",
     "EU 가 2B·Cat 3 를 채택하지 않는 근거, Annex VI 조화분류 목록", "규제"),
    ("R6", "이 파일 · 참고_관할매핑_34행", "이 워크북 내 시트",
     "규제1 이 대상으로 삼는 관할×범주 매핑 원본 34행", "규제"),
    ("R7", "이 파일 · 참고_피부정정_82행", "이 워크북 내 시트",
     "규제2 규칙이 적용된 82셀 상세", "규제"),
    ("R8", "SDS 하이라이트 PDF 폴더",
     "formulation_audit_team_share_highlighted_only_20260630/",
     "데이터1 의 로컬 문서 189개. 시트의 경로를 이 폴더 기준으로 연다", "데이터"),
    ("R9", "감작충돌 32건 장부",
     "04_모델산출물/v6_skinmap_ct/감작충돌_32건_장부.xlsx",
     "정정후보 21건(Annex VI 항목 존재) + 보류 11건(확인 불가). "
     "어떤 라벨도 덮어쓰지 않은 기록용 장부", "데이터"),
    ("R10", "ECHA C&L Inventory / Annex VI 조화분류",
     "ECHA 웹사이트에서 CAS 조회", "R9 의 21건을 대조할 원문", "데이터"),
    ("R11", "감작 CT arm 격자",
     "04_모델산출물/v6_jurisdiction/감작_CTarm_격자.csv · 감작_CTarm_대조.csv",
     "감작 CT arm 4종(A0/A2/A3/A3s)의 성능과 arm 간 paired 대조", "총책임자"),
    ("R12", "전체 지표표",
     "04_모델산출물/v6_jurisdiction/전체지표_ROC_F1.xlsx",
     "ROC-AUC·PR-AUC·F1·MCC·BA·혼동행렬, 임계값 0.5 조건부", "총책임자"),
    ("R13", "관할별 라벨레이어 보고서",
     "04_모델산출물/v6_jurisdiction/보고_관할별_라벨레이어.md",
     "레이어 구조, 관할 분기축, 측정 설계와 한계 전반", "전원"),
    ("R14", "통합 입력 데이터", "04_모델산출물/input_dataset_v6.xlsx",
     "formulation 1675행 / ingredient 5287행. 라벨·출처·문서경로 원본", "데이터"),
    ("R15", "ct_not_applicable 정의부", "01_파이프라인/build_input_v5.py:1168",
     "컬럼은 여기서 만들어지지만 이후 CT 재구성 루프들이 건너뛴다 "
     "(예: attribute_ct_gain_vs_label.py:196)", "모델"),
    ("R16", "parse_tox11", "01_파이프라인/lib_tox11.py:306",
     "현재 근거 3종(눈·피부·감작 서술 텍스트)만 읽는다. "
     "H코드 열 tox_h_statements 를 4번째 근거로 추가하는 것이 D2", "모델"),
    ("R17", "pc2_* 컬럼 순서",
     "01_파이프라인/lib_parse.py:363 NUMERIC_PARAMS · :365 TEXT_PARAMS",
     "둘 다 Python set 리터럴이고 :394~398 루프가 그것을 순회해 컬럼을 만든다. "
     "문자열 해시가 프로세스마다 달라 열 순서가 바뀐다", "모델"),
    ("R18", "회수 A 산출물",
     "04_모델산출물/v6_recovery/회수A_손상21셀_복구.csv · 회수_A_B.xlsx",
     "손상 21셀 복구 기록. 복원된 H코드 19건의 반영 경로가 D2", "모델"),
]
REF_ROWS = [{"자료ID": a, "자료명": b, "경로 / 조항": c, "무엇을 보는가": d,
             "주 사용자": e} for a, b, c, d, e in REFS]
for _, _, p, _, _ in REFS:                    # 내부 경로는 실존을 확인한다
    if p.startswith("04_") or p.startswith("01_") or p.startswith("formul"):
        base = p.split(":")[0].split(" ·")[0].strip()
        assert (ROOT / base).exists(), f"참조 경로 없음: {base}"

# ========================================================== 데이터1 문서 단위
C["doc_rel"] = C["doc_rel"].fillna("").astype(str).str.strip()
C["원문URL"] = C["원문URL"].fillna("").astype(str).str.strip()
C["문서키"] = C["doc_rel"].where(C["doc_rel"] != "", "URL::" + C["원문URL"])
_have = {p: (DOCBASE / p).exists()
         for p in C.loc[C["doc_rel"] != "", "doc_rel"].unique()}
assert all(_have.values()), "실존하지 않는 doc_rel 있음"

EPK = {"eye": "눈", "skin": "피부", "sens": "감작"}
DOCS = []
for key, g in C.groupby("문서키", sort=False):
    local = not key.startswith("URL::")
    row = {"순번": 0, "우선순위": 1 if "eye" in set(g["endpoint"]) else 2,
           "판독대상": " + ".join(EPK[e] for e in ("eye", "skin", "sens")
                              if e in set(g["endpoint"])),
           "문서종류": "로컬 PDF" if local else "웹 URL",
           "참조": "R8" if local else "R14",
           "문서_경로_또는_URL": key if local else key[5:],
           "문서가_덮는_제형수": g["Formulation_ID"].nunique(),
           "product_name": " | ".join(sorted(
               {str(v) for v in g["product_name"].dropna()}))}
    # endpoint 마다 대상 제형을 따로 적는다. 한 문서가 여러 제형을 덮을 때
    # 제형별로 필요한 endpoint 가 다르므로 목록을 공유하면 판정 적용 범위가
    # 넓어진 것처럼 읽힌다.
    for ep in ("eye", "skin", "sens"):
        k = EPK[ep]
        sub = sorted(g.loc[g["endpoint"] == ep, "Formulation_ID"]
                     .astype(str).unique())
        row[f"{k}_판독필요"] = "예" if sub else "—"
        row[f"{k}_대상제형"] = ", ".join(sub) if sub else "—"
        row[f"{k}_현재라벨"] = "NC(음성)" if sub else "—"
        row[f"{k}_학습마스크"] = ("적용중" if ep == "eye" and sub
                             else ("아니오" if sub else "—"))
        row[f"{k}_찾은범주"] = ""
        row[f"{k}_H코드"] = ""
    row.update({"원문근거문장": "", "확인자": "", "확인일": "", "비고": ""})
    DOCS.append(row)
DOCS.sort(key=lambda r: (r["우선순위"], r["문서종류"] != "로컬 PDF",
                         -r["문서가_덮는_제형수"], r["문서_경로_또는_URL"]))
for i, r in enumerate(DOCS, 1):
    r["순번"] = i
assert len(DOCS) == 196

# ========================================================== 규제3 H코드 판정표
HTAB = [
    ("눈", "H318", "Eye Dam. 1", "양성", "양성", "양성", "양성", "확정", "R4 R5", ""),
    ("눈", "H319", "Eye Irrit. 2 / 2A", "양성", "양성", "양성", "양성", "확정",
     "R4 R5", ""),
    ("눈", "H320", "Eye Irrit. 2B", "양성", "음성", "양성", "양성", "미확정",
     "R2 R3 R5", "EU 는 2B 를 채택하지 않는다. K_REACH·US_OSHA 칸은 현재 가정이며 "
                "규제1 의 A1·A4 결과로 확정된다"),
    ("눈", "문구 없음", "해당 없음 → NC", "음성", "음성", "음성", "음성", "확정",
     "", "현재 라벨과 같다는 뜻"),
    ("피부", "H314", "Skin Corr. 1 / 1A / 1B / 1C", "양성", "양성", "양성", "양성",
     "확정", "R4 R5", ""),
    ("피부", "H315", "Skin Irrit. 2", "양성", "양성", "양성", "양성", "확정",
     "R4 R5", ""),
    ("피부", "H316", "Skin Irrit. 3", "양성", "음성", "음성", "음성", "미확정",
     "R2 R3 R4", "UN GHS 만 선택범주로 둔다. K_REACH·US_OSHA 칸은 현재 가정이며 "
                "규제1 의 A2·A5 결과로 확정된다"),
    ("피부", "문구 없음", "해당 없음 → NC", "음성", "음성", "음성", "음성", "확정",
     "", ""),
    ("감작", "H317", "Skin Sens. 1 / 1A / 1B", "양성", "양성", "양성", "양성",
     "확정", "R4 R5", "1A/1B 세분 여부는 A3·A6 이지만 이진 결과는 어느 쪽이든 양성"),
    ("감작", "문구 없음", "해당 없음 → NC", "음성", "음성", "음성", "음성", "확정",
     "", ""),
    ("공통", "판별불가", "Section 11 없음 또는 판독 불가", "학습 제외", "학습 제외",
     "학습 제외", "학습 제외", "확정", "",
     "찾은범주 칸에 판별불가 로 적는다. 추측하지 않는다"),
]
HROWS = [{"endpoint": a, "H코드": b, "GHS 범주": c, "UN_GHS": d, "EU_CLP": e,
          "K_REACH": f_, "US_OSHA": g, "확인상태": h, "참조": i, "비고": j}
         for a, b, c, d, e, f_, g, h, i, j in HTAB]

# ========================================================== 담당별 작업 행
AREF = {"A1": ("R2 R6", "별표 1 — 눈 자극성 항목에 구분 2A / 2B 표기가 있는지. "
                        "구분 2 단일이면 미채택"),
        "A2": ("R2 R6", "별표 1 — 피부 자극성 항목에 구분 3 이 있는지. "
                        "없으면 현재 가정이 맞다"),
        "A3": ("R2 R6", "별표 1 — 피부 과민성이 구분 1 단일인지 1A/1B 세분인지"),
        "A4": ("R3 R6", "Appendix A.3 표 — Category 2A / 2B 구분 표기"),
        "A5": ("R3 R6", "Appendix A.2 표 — Category 3 행 존재 여부"),
        "A6": ("R3 R6", "Appendix A.4 표 — Category 1 단일인지 1A/1B 세분인지")}
규제_A = [{"항목ID": r["항목ID"], "관할": r["관할"], "endpoint": r["endpoint"],
         "확정할 내용": r["결정사항"], "현재 가정": r["현재_가정"],
         "영향 행수": int(r["영향행수"]),
         "참조": AREF[r["항목ID"]][0], "볼 곳": AREF[r["항목ID"]][1],
         "결과": "", "근거 조항": "", "확인자": "", "확인일": ""}
        for _, r in A.iterrows()]

규제_E = [
    {"규칙ID": "E-R1", "정정 내용": "피부 라벨 2A → 2", "해당 셀수": 11,
     "근거": "2A 는 눈 전용 세분범주로, 어느 관할에도 피부 2A 는 없다. "
           "출처는 EPA 피부 II(72h 심한 자극)이고 GHS 로는 Skin Irrit. 2 다.",
     "이진 라벨 변화": "없음", "참조": "R7 R4", "결과": "", "반대 시 대안": "",
     "확인자": "", "확인일": ""},
    {"규칙ID": "E-R2", "정정 내용": "피부 라벨 2B → 2", "해당 셀수": 71,
     "근거": "2B 는 눈 전용 세분범주로, 어느 관할에도 피부 2B 는 없다. "
           "출처는 EPA 피부 III(72h 중등도 자극)이고 GHS 로는 Skin Irrit. 2 다.",
     "이진 라벨 변화": "없음", "참조": "R7 R4", "결과": "", "반대 시 대안": "",
     "확인자": "", "확인일": ""},
]

데이터_2 = [
    {"항목ID": "D3", "할 일": "EyeIrritation6pack_PID81 출처간 모순 1건 판정",
     "상황": "SDS 전사는 EPA 눈 III(24h 내 회복 → NC), NTP 실측은 양성. "
           "현재 라벨 cat_eye_ghs=2B, ntp_epa_eye_cat=3. "
           "두 관할 투영에서 모두 충돌하므로 관할 선택으로 설명되지 않는다",
     "참조": "R14", "볼 곳": "이 제형은 doc_rel·URL 이 모두 결측이다. "
                        "로컬 문서가 없으므로 외부에서 원문을 새로 확보해야 한다",
     "다음 행동": "SDS 또는 NTP 시험보고서 원문 확보 후 어느 쪽이 맞는지 판정",
     "결과": "", "근거": "", "확인자": "", "확인일": ""},
    {"항목ID": "D6", "할 일": "감작 정정후보 21건 ECHA 원문 대조",
     "상황": "장부 32건 중 Annex VI 항목이 존재하는 21건이 정정후보, "
           "확인 불가 11건은 보류. 어떤 라벨도 아직 덮어쓰지 않았다",
     "참조": "R9 R10", "볼 곳": "장부의 cas · 근거URL 열로 ECHA 조회 후 "
                            "판정 · 판정근거 열에 기입",
     "다음 행동": "21건 대조 결과를 장부에 채워 전달",
     "결과": "", "근거": "", "확인자": "", "확인일": ""},
    {"항목ID": "D10", "할 일": "중복 레코드에서 확정라벨을 끌어온 6건 검증",
     "상황": "중복 동일성이 검증되지 않아 적용 보류 상태",
     "참조": "확인 불가", "볼 곳": "이 6건이 어느 산출물에 기록돼 있는지 특정되지 "
                             "않았다. 근거 산출물을 먼저 찾아야 착수할 수 있다",
     "다음 행동": "모델 담당이 근거 산출물을 특정해 전달한 뒤 착수",
     "결과": "", "근거": "", "확인자": "", "확인일": ""},
    {"항목ID": "B1-자료", "할 일": "피부 NC 444행 중 표본의 Draize 평균 점수 수집",
     "상황": "EPA 피부 IV 를 음성 확정으로 쓸 수 있는지가 이 점수로 갈린다. "
           "결정 자체는 총책임자 B1 항목",
     "참조": "R13", "볼 곳": "NTP / EPA 원 시험보고서의 홍반+부종 72h 평균 점수",
     "다음 행동": "표본 20~30건 수집 후 총책임자에 전달",
     "결과": "", "근거": "", "확인자": "", "확인일": ""},
]

MREF = {"D1": ("R15", "정의부를 확인하고, CT 재구성 루프들이 이 컬럼을 건너뛰는 "
                      "지점(attribute_ct_gain_vs_label.py:196 등)에 게이트를 넣는다"),
        "D2": ("R16 R18", "parse_tox11 의 근거 3종에 H코드 열을 4번째로 추가한다. "
                          "복원된 H코드 19건(H317×6 H319×5 H315×4 H318×3 H314×1)이 "
                          "현재 반영 경로가 없다"),
        "D5": ("R14", "composition_mismatch · 중복 제품 · doc_rel 결측 1134행을 "
                      "통합 입력에서 직접 감사한다. 중복 미병합은 폴드 누설 위험"),
        "D9": ("R17", "두 set 리터럴을 순서가 고정된 자료구조로 바꾼다. "
                      "build_input_* 수정 금지 규약과 충돌하므로 예외 승인이 먼저다")}
모델_1 = [{"항목ID": k, "할 일": D.set_index("항목ID").loc[k, "항목"],
         "현재 상태": D.set_index("항목ID").loc[k, "현재상태"],
         "참조": MREF[k][0], "무엇을 하는가": MREF[k][1],
         "승인 필요": D.set_index("항목ID").loc[k, "결정필요"],
         "착수여부": "", "완료 메모": "", "확인자": "", "확인일": ""}
        for k in ("D1", "D2", "D5", "D9")]

BM = B.set_index("항목ID")
DM = D.set_index("항목ID")
총괄 = []
for k in ("B1", "B2", "B3", "B4"):
    총괄.append({"항목ID": k, "구분": "정책",
               "결정할 것": BM.loc[k, "대상"],
               "현재 처리": BM.loc[k, "현재_처리"],
               "결정이 필요한 이유": BM.loc[k, "왜_판단이_필요한가"],
               "선택지": BM.loc[k, "결정선택지"],
               "참조": {"B1": "R13", "B2": "R14",
                      "B3": "R13", "B4": "R13"}[k],
               "자료 수집": {"B1": "데이터 (D3 옆 B1-자료 행)", "B2": "규제",
                         "B3": "모델", "B4": "모델"}[k],
               "결정": "", "결정 근거": "", "결정일": ""})
for k in ("D4", "D7", "D8"):
    총괄.append({"항목ID": k, "구분": DM.loc[k, "구분"].replace("방법론 결정", "방법론"),
               "결정할 것": DM.loc[k, "항목"], "현재 처리": DM.loc[k, "현재상태"],
               "결정이 필요한 이유": DM.loc[k, "왜_중요한가"],
               "선택지": DM.loc[k, "결정필요"],
               "참조": {"D4": "R13", "D7": "R12 R13", "D8": "R12"}[k],
               "자료 수집": "모델", "결정": "", "결정 근거": "", "결정일": ""})
총괄 += [
    {"항목ID": "D11", "구분": "방법론",
     "결정할 것": "감작 CT arm 확정 — A2 통일 / A3 유지 / A3s 교정판",
     "현재 처리": "현행 권고는 A3",
     "결정이 필요한 이유":
         "정규 베이스라인(nan→0)에서 A3 와 A2 의 차이는 +0.0016 이고 유의하지 않다"
         "(t=1.04, 5시드 중 2개만 양수). A3 의 우위는 native 규약에서만 나타난다"
         "(+0.0219, t=7.56, 5/5). 즉 arm 선택이 D7 의 결측 규약 선택에 의존한다",
     "선택지": "① A2 로 통일 — 세 엔드포인트에 같은 규칙, 대가는 native 감작 "
             "AUC −0.0219  ② A3 유지 — native 를 주 규약으로 채택할 경우  "
             "③ A3s 교정판",
     "참조": "R11 R13", "자료 수집": "모델 (측정 완료)",
     "결정": "", "결정 근거": "", "결정일": ""},
    {"항목ID": "D12", "구분": "절차",
     "결정할 것": "이 파일의 판정을 라벨에 반영한 뒤 검증할 담당",
     "현재 처리": "미지정",
     "결정이 필요한 이유": "반영자와 검증자가 같으면 반영 오류를 잡을 수 없다",
     "선택지": "① 규제 담당이 표본 검증  ② 데이터 담당이 표본 검증  "
             "③ 총책임자 직접 승인",
     "참조": "R13", "자료 수집": "—", "결정": "", "결정 근거": "", "결정일": ""},
]

# ========================================================== 안내
GUIDE = [
    ("", ""),
    ("시트 읽는 법", "시트 이름 앞에 담당이 붙어 있다. 자기 담당 시트만 보면 된다. "
                "참고_ 로 시작하는 시트는 근거 자료이고 채울 칸이 없다."),
    ("참조 열", "모든 작업 행에 참조 열이 있고 R1~R18 중 하나를 가리킨다. "
             "참조자료 시트에서 그 ID 의 경로와 조항을 찾을 수 있다. "
             "볼 곳 열에는 그 자료 안에서 어디를 봐야 하는지 적혀 있다."),
    ("", ""),
    ("담당", "시트 / 물량"),
    ("규제", "규제1_관할규정_6건 · 규제2_피부범주정정_2건 · 규제3_H코드_판정표\n"
           "규제1 의 A1·A4(눈 2B, 327행)와 A2·A5(피부 3, 15행)가 라벨을 뒤집는다."),
    ("데이터", "데이터1_원문판독_196문서 · 데이터2_원문확인_4건\n"
            "SDS Section 11 을 열어 H코드와 범주를 적는다. 양성·음성은 적지 않는다."),
    ("모델", "모델1_코드작업_4건"),
    ("총책임자", "총책임자_결정_9건"),
    ("", ""),
    ("데이터 담당이 먼저 볼 것",
     "규제3_H코드_판정표. 문서에서 찾을 것은 H코드이고 그것이 관할별로 양성인지 "
     "음성인지는 그 표가 정한다. 그래서 규제1 의 A1·A4 가 아직 미결이어도 원문 "
     "판독을 지금 시작할 수 있다. A1·A4 가 확정되면 판정표 두 줄만 갱신하면 되고 "
     "판독 결과는 다시 보지 않는다."),
    ("작업 순서", "1. 데이터1 우선순위 1 (눈 포함 97문서). 눈 109행은 이미 학습에서 "
              "마스크된 상태여서 판독 결과가 바로 반영된다.\n"
              "2. 규제1·규제2 는 데이터1 과 동시에 진행할 수 있다.\n"
              "3. 데이터1 우선순위 2 (99문서, 피부·감작)는 총책임자 B4 결정 뒤에 "
              "착수한다. B4 가 현재대로 유지로 결정되면 이 판독은 라벨에 반영되지 "
              "않는다."),
    ("", ""),
    ("입력 규칙", ""),
    ("1", "연노랑 칸만 채운다. 연회색 칸은 읽기 전용이다."),
    ("2", "확실하지 않으면 비워 두거나 확인 불가 로 적는다. 추측으로 채우지 않는다. "
         "빈 칸은 미완으로 집계된다."),
    ("3", "판정에는 근거를 같이 적는다. 조항 번호 또는 원문 문장."),
    ("4", "이 파일은 라벨을 직접 바꾸지 않는다. 원본 라벨은 보존되고 정정은 별도 "
         "레이어로 기록된다."),
]

# ========================================================== 조립
wb = openpyxl.Workbook()
wb.remove(wb.active)

ws = wb.create_sheet("0_먼저읽기")
ws.cell(1, 1, WB_OUT.stem).font = Font(bold=True, size=14, color="1F4E79")
for i, (a, b) in enumerate(GUIDE, 3):
    ca, cb = ws.cell(i, 1, a), ws.cell(i, 2, b)
    ca.font = Font(bold=True, size=10)
    ca.alignment = Alignment(vertical="top", wrap_text=True)
    cb.alignment = Alignment(vertical="top", wrap_text=True)
    if a in ("담당", "입력 규칙"):
        ca.fill = cb.fill = PatternFill("solid", fgColor="D9E2F3")
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 104
ws.sheet_view.showGridLines = False

sheet(wb, "참조자료", ["자료ID", "자료명", "경로 / 조항", "무엇을 보는가", "주 사용자"],
      REF_ROWS, H["참고"], fill_cols=(),
      widths={"자료ID": 8, "자료명": 40, "경로 / 조항": 56, "무엇을 보는가": 52,
              "주 사용자": 12},
      note="각 작업 시트의 참조 열이 이 표의 자료ID 를 가리킨다. "
           "내부 경로는 프로젝트 최상위 폴더 기준이다.",
      wrap_cols=("자료명", "경로 / 조항", "무엇을 보는가"))

sheet(wb, "규제1_관할규정_6건",
      ["항목ID", "관할", "endpoint", "확정할 내용", "현재 가정", "영향 행수",
       "참조", "볼 곳", "결과", "근거 조항", "확인자", "확인일"],
      규제_A, H["규제"],
      fill_cols=("결과", "근거 조항", "확인자", "확인일"),
      widths={"확정할 내용": 42, "현재 가정": 26, "볼 곳": 50, "결과": 14,
              "근거 조항": 28, "항목ID": 8, "관할": 10, "endpoint": 9,
              "영향 행수": 9, "참조": 10},
      note="이 6건이 라벨 정의를 확정한다. A1·A4 는 눈 2B(327행), "
           "A2·A5 는 피부 3(15행)에 영향을 준다.",
      dv=[(("결과",), ("채택", "미채택", "확인 불가"))],
      wrap_cols=("확정할 내용", "현재 가정", "볼 곳"))

sheet(wb, "규제2_피부범주정정_2건",
      ["규칙ID", "정정 내용", "해당 셀수", "근거", "이진 라벨 변화", "참조",
       "결과", "반대 시 대안", "확인자", "확인일"],
      규제_E, H["규제"],
      fill_cols=("결과", "반대 시 대안", "확인자", "확인일"),
      widths={"정정 내용": 22, "근거": 66, "이진 라벨 변화": 14, "결과": 14,
              "반대 시 대안": 28, "규칙ID": 9, "해당 셀수": 10, "참조": 10},
      note="82셀에 적용된 정정을 규칙 2건으로 승인한다. 셀별 상세는 "
           "참고_피부정정_82행 시트에 있다.",
      dv=[(("결과",), ("승인", "반대", "확인 불가"))],
      wrap_cols=("근거", "정정 내용"))

sheet(wb, "규제3_H코드_판정표",
      ["endpoint", "H코드", "GHS 범주", "UN_GHS", "EU_CLP", "K_REACH", "US_OSHA",
       "확인상태", "참조", "비고"],
      HROWS, H["규제"], fill_cols=(),
      widths={"GHS 범주": 30, "비고": 60, "endpoint": 9, "H코드": 11,
              "확인상태": 10, "참조": 11},
      note="데이터 담당이 판독할 때 참조하는 표다. 미확정 두 줄(H320, H316)의 "
           "K_REACH·US_OSHA 칸은 현재 가정이며 규제1 결과로 확정된다.",
      wrap_cols=("GHS 범주", "비고"))

sheet(wb, "데이터1_원문판독_196문서",
      ["순번", "우선순위", "판독대상", "문서종류", "참조", "문서_경로_또는_URL",
       "문서가_덮는_제형수", "product_name",
       "눈_판독필요", "눈_대상제형", "눈_현재라벨", "눈_학습마스크", "눈_찾은범주",
       "눈_H코드",
       "피부_판독필요", "피부_대상제형", "피부_현재라벨", "피부_학습마스크",
       "피부_찾은범주", "피부_H코드",
       "감작_판독필요", "감작_대상제형", "감작_현재라벨", "감작_학습마스크",
       "감작_찾은범주", "감작_H코드",
       "원문근거문장", "확인자", "확인일", "비고"],
      DOCS, H["데이터"],
      fill_cols=("눈_찾은범주", "눈_H코드", "피부_찾은범주", "피부_H코드",
                 "감작_찾은범주", "감작_H코드", "원문근거문장", "확인자", "확인일",
                 "비고"),
      widths={"문서_경로_또는_URL": 66, "product_name": 28, "원문근거문장": 44,
              "순번": 6, "우선순위": 8, "판독대상": 15, "문서종류": 10,
              "문서가_덮는_제형수": 9, "비고": 18, "참조": 8,
              "눈_대상제형": 24, "피부_대상제형": 24, "감작_대상제형": 24},
      note="문서 1개가 1행이다. 참조 R8 은 로컬 PDF 폴더이고 경로 열을 그 폴더 "
           "기준으로 연다. Section 11 을 보고 찾은범주와 H코드를 적는다. "
           "양성·음성은 적지 않는다 — 규제3_H코드_판정표(R1)가 변환한다. "
           "판독필요가 — 인 endpoint 는 이 문서의 대상이 아니고, 판정은 각 "
           "endpoint 의 대상제형에만 적용된다.",
      dv=[(("눈_찾은범주",),
           ("Eye Dam. 1", "Eye Irrit. 2", "Eye Irrit. 2A", "Eye Irrit. 2B",
            "문구 없음(NC)", "판별불가")),
          (("피부_찾은범주",),
           ("Skin Corr. 1", "Skin Corr. 1A", "Skin Corr. 1B", "Skin Corr. 1C",
            "Skin Irrit. 2", "Skin Irrit. 3", "문구 없음(NC)", "판별불가")),
          (("감작_찾은범주",),
           ("Skin Sens. 1", "Skin Sens. 1A", "Skin Sens. 1B", "문구 없음(NC)",
            "판별불가"))],
      wrap_cols=("문서_경로_또는_URL", "눈_대상제형", "피부_대상제형",
                 "감작_대상제형", "product_name"))

sheet(wb, "데이터2_원문확인_4건",
      ["항목ID", "할 일", "상황", "참조", "볼 곳", "다음 행동", "결과", "근거",
       "확인자", "확인일"],
      데이터_2, H["데이터"],
      fill_cols=("결과", "근거", "확인자", "확인일"),
      widths={"할 일": 38, "상황": 52, "볼 곳": 44, "다음 행동": 34, "결과": 24,
              "근거": 28, "항목ID": 10, "참조": 10},
      wrap_cols=("할 일", "상황", "볼 곳", "다음 행동"))

sheet(wb, "모델1_코드작업_4건",
      ["항목ID", "할 일", "현재 상태", "참조", "무엇을 하는가", "승인 필요",
       "착수여부", "완료 메모", "확인자", "확인일"],
      모델_1, H["모델"],
      fill_cols=("착수여부", "완료 메모", "확인자", "확인일"),
      widths={"할 일": 42, "현재 상태": 44, "무엇을 하는가": 50, "승인 필요": 30,
              "착수여부": 10, "완료 메모": 28, "항목ID": 9, "참조": 10},
      dv=[(("착수여부",), ("착수", "보류", "완료"))],
      wrap_cols=("할 일", "현재 상태", "무엇을 하는가", "승인 필요"))

sheet(wb, "총책임자_결정_9건",
      ["항목ID", "구분", "결정할 것", "현재 처리", "결정이 필요한 이유", "선택지",
       "참조", "자료 수집", "결정", "결정 근거", "결정일"],
      총괄, H["총괄"],
      fill_cols=("결정", "결정 근거", "결정일"),
      widths={"결정할 것": 40, "현재 처리": 36, "결정이 필요한 이유": 56,
              "선택지": 52, "자료 수집": 18, "결정": 20, "결정 근거": 32,
              "항목ID": 9, "구분": 10, "참조": 10},
      note="다른 담당이 기다리는 항목이 있다. B4 는 데이터1 우선순위 2(99문서)의 "
           "착수 여부를 정하고, D7 과 D11 은 한 쌍으로 감작 CT arm 을 정한다.",
      wrap_cols=("결정할 것", "현재 처리", "결정이 필요한 이유", "선택지"))

sheet(wb, "참고_관할매핑_34행", list(F.columns), F.to_dict("records"), H["참고"],
      fill_cols=(), widths={"근거": 88, "정본범주": 11, "처리": 20},
      note="규제1 이 대상으로 삼는 관할×범주 매핑 원본. 채울 칸은 없다.",
      wrap_cols=("근거",))

sheet(wb, "참고_피부정정_82행", list(E.columns), E.to_dict("records"), H["참고"],
      fill_cols=(), widths={"정정근거": 78, "Formulation_ID": 22},
      note="규제2 의 두 규칙이 적용된 82셀. 승인은 규제2 시트에서 한다.",
      wrap_cols=("정정근거",))

wb.save(WB_OUT)
print(f"저장 → {WB_OUT}")
print("시트:", ", ".join(wb.sheetnames))
print(f"참조자료 {len(REF_ROWS)}건 / 규제 {len(규제_A)}+{len(규제_E)} / "
      f"데이터 {len(DOCS)}문서+{len(데이터_2)} / 모델 {len(모델_1)} / "
      f"총책임자 {len(총괄)}")
