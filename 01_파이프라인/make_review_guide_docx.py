#!/usr/bin/env python3
"""수동검토_최종판단_작업파일.xlsx 를 쉽게 이해하도록 안내하는 docx 생성.

작업파일 자체는 건드리지 않는다(읽기만 한다). 이 스크립트는 그 파일의 시트
구조·건수·입력 방법을 설명하는 별도 문서를 만든다. 수치는 실제 워크북에서
읽어 확인하고, 추측으로 채우지 않는다.
"""
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path("/Users/hanseoyun/Desktop/260830")
WB_PATH = ROOT / "04_모델산출물" / "v6_수동검토" / "수동검토_최종판단_작업파일.xlsx"
OUT = ROOT / "04_모델산출물" / "v6_수동검토" / "작업파일_안내.docx"

wb = openpyxl.load_workbook(WB_PATH, read_only=True)
COUNT = {ws.title: ws.max_row - 1 for ws in wb.worksheets}
assert COUNT["A_관할규정_확정"] == 6
assert COUNT["B_행단위_판별불가"] == 4
assert COUNT["C_negrec_원문재수집"] == 448
assert COUNT["D_기타_미완항목"] == 10
assert COUNT["E_L1정정_82셀"] == 82
assert COUNT["F_부분확인_레지스트리"] == 34

doc = Document()
for name in ("Normal",):
    st = doc.styles[name]
    st.font.name = "맑은 고딕"
    st.font.size = Pt(10.5)

H1 = RGBColor(0x1F, 0x3A, 0x5F)


def h1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = H1
    return p


def h2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = H1
    return p


def para(text, bold=False, italic=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    return t


# ------------------------------------------------------------------ 표지
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("수동검토 · 최종판단 작업파일 — 이해 안내서")
r.bold = True
r.font.size = Pt(20)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("대상 파일: 04_모델산출물/v6_수동검토/수동검토_최종판단_작업파일.xlsx")
r.italic = True
r.font.size = Pt(11)
doc.add_paragraph()

para("이 문서가 하는 일", bold=True)
para(
    "작업파일 자체를 대체하지 않습니다. 시트 7개가 각각 무엇을 위한 것인지, "
    "누가 무엇을 채워야 하는지, 채운 뒤 어떻게 반영되는지를 사람이 읽기 쉽게 "
    "정리한 것입니다. 실제 작업은 xlsx 파일에서 합니다."
)
para(
    "핵심 원칙: 노란 칸(입력란)만 채웁니다. 확실하지 않으면 비워 두거나 "
    "'확인 불가'라고 적습니다 — 추측으로 채우지 않습니다. 이 파일은 라벨을 "
    "직접 바꾸지 않습니다. 원본 라벨(L0)은 어떤 경우에도 보존됩니다.",
    bold=True,
)

doc.add_page_break()

# ------------------------------------------------------------------ 큰 그림
h1("1. 큰 그림 — 이 파일이 왜 필요한가")
para(
    "모델은 '기업이 신고한 SDS 라벨'을 그대로 정답으로 쓸 수 없는 지점들을 "
    "만났습니다. 예를 들어 규제 조문을 재확인해야 하는 경우, 서로 다른 원문 "
    "출처의 판정이 어긋나는 경우, 특정 값이 두 범주에 걸쳐 있어 자동으로는 "
    "확정할 수 없는 경우입니다. 이런 판단은 코드가 대신할 수 없고, 사람이 "
    "원문을 보고 정해야 합니다. 그 판단이 필요한 항목을 모델 코드에서 "
    "분리해 이 파일 한 곳에 모았습니다."
)
para("이 파일은 다음 순서로 보면 됩니다.")
bullet("0_작업안내 — 시트 전체 요약표 (이미 이 문서가 더 자세히 설명합니다)")
bullet("A_관할규정_확정 — 가장 먼저 볼 것. 라벨을 실제로 뒤집는 4건이 있습니다")
bullet("C_negrec_원문재수집 — 가장 분량이 많음(448행). 우선순위가 매겨져 있습니다")
bullet("B_행단위_판별불가, D_기타_미완항목 — 정책·방법론 결정")
bullet("E_L1정정_82셀 — 이미 적용된 자동 정정에 대한 승인/반대만 표시")
bullet("F_부분확인_레지스트리 — 입력 없음, 추적용으로만 봅니다")

doc.add_page_break()

# ------------------------------------------------------------------ 시트별 안내
h1("2. 시트별 안내")

h2("A_관할규정_확정 — 6행 · 규제 담당자")
para(
    "무엇: 한국(K-REACH)과 미국(US-OSHA)이 GHS 의 선택적 세부범주를 실제로 "
    "채택했는지가 아직 '가정'상태입니다. UN GHS 와 EU CLP 는 원문 확인이 "
    "끝났지만(=확인), K-REACH·US-OSHA 는 아직입니다(=부분확인)."
)
para("확인할 두 가지 축:")
bullet("Eye Irrit. 2B (H320) 채택 여부 — 채택 안 하면 눈 라벨 327행이 양성→음성으로 바뀝니다")
bullet(
    "Skin Irrit. 3 (H316) 채택 여부 — 채택하면 피부 라벨 15행이 음성→양성으로 "
    "바뀌고, 동시에 EPA 자극성 IV 유래 444행이 '판별불가'로 마스크됩니다"
)
para(
    "해야 할 일: 각국 고시(화학물질안전원 GHS 고시, OSHA HCS 2012)의 부속서 "
    "원문을 찾아 해당 세부범주가 채택 목록에 있는지 확인하고, '확인결과' 열에 "
    "채택/미채택을 적습니다. 조항 번호를 '근거조항' 열에 반드시 같이 적습니다."
)
para("영향의 크기 (라벨을 실제로 뒤집는 4건):", bold=True)
table(
    ["ID", "관할", "결정", "현재 가정", "뒤집힐 때"],
    [
        ["A1", "K-REACH", "Eye 2B 채택?", "채택", "눈 327행 양성→음성"],
        ["A2", "K-REACH", "Skin Cat 3 채택?", "미채택", "피부 15행 음성→양성 + EPA IV 444행 마스크"],
        ["A4", "US-OSHA", "Eye 2B 채택?", "채택", "눈 327행 양성→음성"],
        ["A5", "US-OSHA", "Skin Cat 3 채택?", "미채택", "피부 15행 + EPA IV 444행"],
    ],
)
para(
    "나머지 2건(A3·A6, 감작 1A/1B 세분 여부)은 이진 분류에는 영향이 없습니다 "
    "— 나중에 해도 됩니다.",
    italic=True,
)

h2("B_행단위_판별불가 — 4행 · 총책임자")
para(
    "무엇: 자동 규칙으로는 양성/음성을 정할 수 없는 상황 4가지입니다. 예: 공급자 "
    "SDS 가 '분류되지 않음'이라고만 적었는데, 그 공급자의 관할이 특정 범주를 "
    "채택하지 않는 곳이면 그 침묵이 정말 '무위험'인지 '그 범주를 안 쓰는 지역이라 "
    "말 안 한 것'인지 구분이 안 됩니다."
)
para(
    "해야 할 일: 각 행에 제시된 선택지 중 하나를 정책적으로 고릅니다(원문 대조가 "
    "아니라 판단 사안입니다). '판단결과'와 그 이유를 적습니다."
)

h2("C_negrec_원문재수집 — 448행 · 데이터 담당자 (분량 가장 큼)")
para(
    "무엇: SDS 원문에서 회수된 뒤 자동으로 '음성(NC)'으로 처리된 셀들입니다. "
    "회수 당시 원문을 다시 대조하지 못했기 때문에, 정말 음성이 맞는지 "
    "한 건씩 원문을 봐야 합니다. 특히 눈(eye) 109행은 다른 출처의 음성 값들과 "
    "비교했을 때 통계적으로 수상합니다(양성 쪽 패턴에 가까움) — 그래서 이 "
    "109행은 모델 학습에서 이미 빼놓았고(마스크), 우선순위 1로 맨 위에 옵니다."
)
para("실행 가능성부터 확인하십시오:", bold=True)
table(
    ["대상", "행 수", "URL 있음", "로컬 문서 있음"],
    [
        ["눈 (우선순위 1, 학습 마스크 적용됨)", "109", "56", "1"],
        ["피부 + 감작 (우선순위 2, 마스크 안 함)", "339", "143", "4"],
        ["합계", "448", "199", "5"],
    ],
)
para(
    "즉 눈 109행 중 53행은 URL 자체가 데이터에 없어 제품명으로 SDS 를 새로 "
    "찾아야 합니다. 로컬 문서가 있는 행(doc_source/doc_rel 열 참고)부터 처리하면 "
    "즉시 착수할 수 있고, 정렬이 그 순서로 되어 있습니다."
)
para(
    "해야 할 일: 각 행의 원문(URL 또는 로컬 파일)에서 Section 11(독성 정보)을 "
    "직접 읽고 '판정결과' 열에 양성/음성/판별불가 중 하나를 적습니다. 근거가 된 "
    "H코드와 문장을 그대로(창작 없이) '원문근거문장'에 옮겨 적습니다."
)

h2("D_기타_미완항목 — 10행 · 혼합(코드/정책/재지시)")
para(
    "무엇: 코드로 구현이 안 끝난 것, 방법론 결정이 필요한 것, 팀원 재지시가 "
    "필요한 것들을 모아둔 잡동사니 시트입니다. 예: 결측값을 0 으로 볼지 "
    "'미측정'으로 남길지, 임계값을 0.5 에서 바꿀지, 김보경 담당 274건 pH 자료가 "
    "아직 전혀 도착하지 않은 상태 등입니다."
)
para(
    "해야 할 일: '누가' 열에 적힌 담당(총책임자/데이터 담당/코드 담당)이 "
    "'결정필요' 항목을 검토하고 '판단결과'를 적습니다. 코드 구현이 필요한 항목은 "
    "판단만 적어두면 이후 반영합니다."
)

h2("E_L1정정_82셀 — 82행 · 규제 담당자 (승인/반대만)")
para(
    "무엇: 피부 자극 라벨의 '2A'와 '2B' 표기를 GHS 정본 범주인 '2'로 통일한 "
    "정정 82건입니다. 이미 적용되어 모델에 반영돼 있습니다 — 이 시트는 그 "
    "정정이 맞는지 사후 확인용입니다. 눈 자극에서만 쓰는 2A/2B 세분이 피부 "
    "표기에 잘못 섞여 들어온 것이므로 대부분 명백한 오탈자 정정입니다."
)
para("해야 할 일: 각 행을 훑어보고 이상이 있으면 '확인결과'에 반대라고 표시합니다.")

h2("F_부분확인_레지스트리 — 34행 · 입력 없음")
para(
    "무엇: 관할별 매핑 규칙 중 아직 '확인' 단계에 못 이른 항목의 목록입니다. "
    "A 시트가 채워지면 이 레지스트리의 확인상태가 자동으로 바뀝니다. 지금은 "
    "그냥 무엇이 '가정'에 기반해 있는지 훑어보는 용도입니다."
)

doc.add_page_break()

# ------------------------------------------------------------------ 입력규칙
h1("3. 공통 입력 규칙")
bullet("확실하지 않으면 비워 두거나 '확인 불가'라고 적습니다. 추측으로 채우지 않습니다.")
bullet("판정에는 반드시 근거(조항 번호 또는 원문 문장)를 같이 적습니다.")
bullet("이 파일은 라벨을 직접 바꾸지 않습니다. 원본 라벨(L0)은 항상 보존됩니다.")
bullet("노란 칸(입력란)만 채웁니다. 다른 칸은 참고용으로 이미 채워져 있습니다.")
bullet("작업이 끝난 시트는 알려주시면 반영하겠습니다 — 모델 코드는 이 파일을 자동으로 읽지 않습니다.")

para("")
para(
    f"생성: {WB_PATH.name} 기준, 시트 7개 · 건수는 워크북에서 직접 읽어 확인함 "
    "(A 6 / B 4 / C 448 / D 10 / E 82 / F 34).",
    italic=True,
    size=9,
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("saved", OUT)
