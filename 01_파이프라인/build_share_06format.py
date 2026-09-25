"""점검본(`07_화학연구원_공유_260918/`)에서 **우리 검수 흔적을 모두 걷어내고**
`06_화학연구원_공유_260916/` 과 같은 형식으로 배포본을 만든다.

06 판의 형식 그대로다.

  1. 워크북 하나 `제품코드별_성분리스트.xlsx` — 시트는 `제형`·`성분` 둘.
     첫 행 틀고정과 열너비도 06 워크북에서 읽어 그대로 입힌다
  2. 문서는 **제형별 폴더**에 두고 파일 이름은 사람이 읽는 형태다 —
     `MSDS_하이라이트사본/<제품코드>/<제품코드>__<순번>__<도메인>.pdf`,
     `대체근거_참고/<제품코드>/<제품코드>__surrogate<n>__…`
  3. 행과 파일의 연결은 `근거매니페스트.csv` — 06 과 달리 판정 열(`등급`·
     `증거_신뢰도`)과 감사 경로(`원경로`)는 빼고 제품코드·종류·경로·sha256·크기만

**입력은 점검본의 워크북과 문서 폴더 둘뿐이다.** 점검용 csv 는 한 개도 읽지
않는다 — 점검본에서 csv·로그를 지워도 이 배포본은 그대로 다시 만들어진다.

**걷어낸 검수 흔적.** `요약` 시트(설명), `데이터_이슈` 시트, 제형 시트의
`MSDS_확보상태`·`증거_신뢰도`·`재수집후보_URL`, 성분 시트의 원문대조·검토필요·정정
후보·정규화·농도환산 관련 열 전부, README·PACKAGE_VERSION.json·생성로그.
남는 열은 **전부 v6 원값**이다.

**06 판과 다른 점 셋.**

  - 사본이 06 의 옛 단색 노랑 마킹본이 아니라 **07 의 3색 재도색본**이다.
    성분명 초록·CAS 노랑·농도 파랑으로 칠하고 주석에 「원문 글자 → 공유값」을
    병기했다. 값을 원문까지 되짚는 유일한 수단이라 남긴다
  - 제품 문서가 없던 제형을 다시 찾아 확보한 문서 5건이 더해져 703 → 708건이다.
    이 5건은 주석이 없는 원본 PDF 다
  - 제형 시트에 `ph_best`(원액 기준), 성분 시트에 `ing_idx`·`ing_pct_kind_best`
    를 더했다. `ing_idx` 는 하이라이트 주석의 `[제품코드 #행번호]` 를 성분 행에
    맞춰 보는 데 필요하다

    python3 build_share_06format.py
"""
from __future__ import annotations

import hashlib
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import lib_model as L                                            # noqa: E402
from build_share_dist import relabel                             # noqa: E402

SRC = L.ROOT / "07_화학연구원_공유_260918"
REF = L.ROOT / "06_화학연구원_공유_260916"
OUT = L.ROOT / "10_화학연구원_배포_260920"
XLSX = "제품코드별_성분리스트.xlsx"
MANI = "근거매니페스트.csv"
# 점검본 폴더 이름 → 배포본 폴더 이름 · 매니페스트 `종류`
DOCDIRS = {"MSDS_하이라이트": ("MSDS_하이라이트사본", "제품단위"),
           "대체근거_참고": ("대체근거_참고", "대체근거")}
log = L.make_logger(L.ROOT / "01_파이프라인" / "배포본_06형식.log")

# 06 판 시트의 열 순서를 따르고, v6 원값인 pH·행번호·농도단위만 더한다
FORM_COLS = ["Formulation_ID", "product_name", "formulation_type_ko",
             "y_eye", "y_skin", "y_sens", "ph_best"]
ING_COLS = ["Formulation_ID", "ing_idx", "ing_name_best", "ing_cas_best",
            "ing_pct_best", "ing_pct_kind_best", "ing_role_final",
            "ing_ghs_eye", "ing_ghs_skin", "ing_ghs_sens",
            "ing_ghs_indep_eye_cat", "ing_ghs_indep_eye_tier",
            "ing_ghs_indep_skin_cat", "ing_ghs_indep_skin_tier",
            "ing_ghs_indep_sens_cat", "ing_ghs_indep_sens_tier", "ing_src"]
MANI_COLS = ["Formulation_ID", "종류", "경로", "sha256", "size_bytes"]
# 06 판에 없던 열의 너비만 여기서 정한다. 나머지는 06 워크북에서 읽어 온다
NEW_WIDTH = {"ph_best": 10.0, "ing_idx": 10.0, "ing_pct_kind_best": 19.0}
# 배포본 어디에도 남으면 안 되는 판정 어휘. `정정` 은 제외한다 — MIX408 성분명에
# 팀원이 남긴 v6 원값 메모(`98%→25%로 정정…`)가 들어 있고 그건 검수 흔적이 아니다
FORBID = ["검토필요", "원문확인", "원문대조", "재수집", "확보상태", "신뢰도",
          "EVIDENCE_WEAK", "미검증", "규칙일치", "좌표검증", "정규화", "이전값",
          "판정", "표기변형"]


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def surrogate_name(name: str, n_in_dir: int) -> str:
    """대체근거 파일 이름의 순번 자리를 06 판 표기로 되돌린다.

    제품단위는 두 판이 같은 이름이고(`<제품코드>__<순번>__<도메인>`), 대체근거만
    06 은 `__surrogate__`(폴더에 한 건일 때)·`__surrogate<n>__`(여럿일 때)로 적었다.
    폴더 안 파일 수만 보면 되므로 점검용 csv 가 필요 없다."""
    fid, seq, rest = name.split("__", 2)
    return f"{fid}__surrogate{'' if n_in_dir == 1 else seq}__{rest}"


def copy_docs() -> list[dict]:
    """점검본 문서 폴더를 옮긴다. 제형별 폴더 구조는 두 판이 이미 같다."""
    rows = []
    for src_dir, (dst_dir, kind) in DOCDIRS.items():
        n_form = 0
        for d in sorted((SRC / src_dir).iterdir()):
            if not d.is_dir():
                continue
            n_form += 1
            (OUT / dst_dir / d.name).mkdir(parents=True, exist_ok=True)
            docs = [p for p in sorted(d.iterdir()) if p.is_file()
                    and p.suffix.lower() in (".pdf", ".html", ".htm")]
            for p in docs:
                name = p.name if kind == "제품단위" else surrogate_name(p.name, len(docs))
                rel = f"{dst_dir}/{d.name}/{name}"
                shutil.copy2(p, OUT / rel)
                rows.append({"Formulation_ID": d.name, "종류": kind, "경로": rel})
            if len(docs) != len({r["경로"] for r in rows if r["Formulation_ID"] == d.name
                                 and r["종류"] == kind}):
                raise SystemExit(f"이름이 겹쳐 덮어썼다: {d.name}")
        log(f"{dst_dir}/ — 제형 {n_form:,}개 · "
            f"파일 {sum(1 for r in rows if r['종류'] == kind):,}건")
    return rows


def style(path: Path) -> None:
    """06 워크북의 보기 양식을 그대로 입힌다 — 첫 행 틀고정 · 열너비 · 머리글 굵게 없음."""
    ref = load_workbook(REF / XLSX)
    wid = {}
    # 06 의 `데이터_이슈` 시트에 같은 이름의 열이 있어 너비가 덮어써진다. 두 시트만 읽는다
    for s in ("제형", "성분"):
        ws = ref[s]
        for i, c in enumerate(next(ws.iter_rows(max_row=1, values_only=True))):
            d = ws.column_dimensions.get(get_column_letter(i + 1))
            if c and d and d.width:
                wid[str(c)] = d.width
    ref.close()
    wid |= NEW_WIDTH
    w = load_workbook(path)
    for ws in w.worksheets:
        ws.freeze_panes = "A2"
        for i, c in enumerate(next(ws.iter_rows(max_row=1, values_only=True))):
            ws.cell(1, i + 1).font = Font(bold=False)
            ws.column_dimensions[get_column_letter(i + 1)].width = wid[str(c)]
    w.save(path)


def main() -> None:
    BOOK = SRC / "02_제품코드별_성분리스트.xlsx"
    FORM = pd.read_excel(BOOK, sheet_name="제형")
    ING = pd.read_excel(BOOK, sheet_name="성분")
    for name, d, cols in (("제형", FORM, FORM_COLS), ("성분", ING, ING_COLS)):
        miss = [c for c in cols if c not in d.columns]
        if miss:
            raise SystemExit(f"{name} 시트에 없는 열: {miss}")
        log(f"{name}: {len(d):,}행 · 남긴 열 {len(cols)} · 뺀 열 {len(d.columns) - len(cols)}")

    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)

    rows = copy_docs()
    n_pdf, n_annot = relabel(OUT, ING)
    log(f"정규화 주석 문구 교체 — 문서 {n_pdf}건 · 주석 {n_annot}개")

    for r in rows:                      # 재도색 뒤의 사본으로 해시를 잡는다
        p = OUT / r["경로"]
        r["sha256"], r["size_bytes"] = sha(p), p.stat().st_size
    with pd.ExcelWriter(OUT / XLSX, engine="openpyxl") as w:
        FORM[FORM_COLS].to_excel(w, sheet_name="제형", index=False)
        ING[ING_COLS].to_excel(w, sheet_name="성분", index=False)
    style(OUT / XLSX)
    pd.DataFrame(rows)[MANI_COLS].to_csv(OUT / MANI, index=False)

    # ---- 검사. 하나라도 어긋나면 배포하지 않는다 ----
    X = pd.ExcelFile(OUT / XLSX)
    if X.sheet_names != ["제형", "성분"]:
        raise SystemExit(f"시트 구성이 다르다: {X.sheet_names}")
    F2, I2 = (pd.read_excel(OUT / XLSX, sheet_name=s) for s in ("제형", "성분"))
    if (len(F2), len(I2)) != (len(FORM), len(ING)):
        raise SystemExit(f"행수가 어긋난다 {len(F2)}·{len(I2)}")
    bad_ext = sorted({p.suffix.lower() for p in OUT.rglob("*") if p.is_file()}
                     - {".pdf", ".html", ".htm", ".xlsx", ".csv"})
    if bad_ext:
        raise SystemExit(f"있어서는 안 되는 확장자: {bad_ext}")
    mine = {r["경로"] for r in rows}
    ondisk = {p.relative_to(OUT).as_posix() for p in OUT.rglob("*")
              if p.is_file() and p.suffix.lower() not in (".xlsx", ".csv")}
    if ondisk != mine:
        raise SystemExit(f"매니페스트에 없는 문서 {len(ondisk - mine)}건 · "
                         f"없는 문서를 가리킴 {len(mine - ondisk)}건")
    # 문서 폴더 이름이 제형 시트에 없으면 받는 쪽이 행을 못 찾는다
    orphan = sorted({r["Formulation_ID"] for r in rows} - set(F2.Formulation_ID))
    if orphan:
        raise SystemExit(f"제형 시트에 없는 문서 폴더 {len(orphan)}개: {orphan[:3]}")
    # 06 판이 쓰던 이름이 하나도 빠지지 않았는가 (06 폴더가 남아 있을 때만)
    if (REF / MANI).exists():
        lost = sorted(set(pd.read_csv(REF / MANI).경로.astype(str)) - mine)
        if lost:
            raise SystemExit(f"06 이름이 빠졌다 {len(lost)}건: {lost[:3]}")
    # 판정 어휘가 시트·매니페스트·HTML 사본에 남아 있지 않은가
    hit = {}
    for name, d in (("제형", F2), ("성분", I2), (MANI, pd.read_csv(OUT / MANI))):
        t = "\n".join(list(d.columns) + d.astype(str).to_numpy().ravel().tolist())
        for word in FORBID:
            if word in t:
                hit[f"{name}:{word}"] = t.count(word)
    for p in OUT.rglob("*.htm*"):
        t = p.read_text("utf-8", "ignore")
        for word in ("정규화", "이전값"):
            if word in t:
                hit[f"{p.name}:{word}"] = t.count(word)
    if hit:
        raise SystemExit(f"판정 어휘가 남았다: {hit}")

    log("")
    log(f"제형 {len(F2):,}행 {len(F2.columns)}열 · 성분 {len(I2):,}행 "
        f"{len(I2.columns)}열 · 근거 {len(rows):,}건")
    log(f"완료 → {OUT.name}/ (파일 {sum(1 for p in OUT.rglob('*') if p.is_file()):,}건 · "
        f"{sum(p.stat().st_size for p in OUT.rglob('*') if p.is_file()) / 1024 ** 2:.0f} MB)")
    print(f"→ {OUT.name}")


if __name__ == "__main__":
    main()
